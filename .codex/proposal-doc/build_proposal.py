from __future__ import annotations

import hashlib
import shutil
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
REFERENCE = Path(r"C:\Users\Windows\Downloads\23520286_23520434_Detailed_Project_Proposal_DA2.docx")
OUTPUT = ROOT / "UIT_Career_Hub_Detailed_Project_Proposal.docx"
DIAGRAMS = Path(__file__).resolve().parent / "diagrams"
EXPECTED_SHA256 = "86E3E201AEAC9E7B5505368A2F5AB9A975E7A18F32FE2BF2A26075E580D679BF"
NUMBERED_NUM_ID = 1


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def set_run_font(run, size=13, bold=None, italic=None):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{key}"), "Times New Roman")


def format_paragraph(
    paragraph,
    *,
    alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
    left=0,
    first=None,
    before=0,
    after=2,
    line=1.3,
    keep_next=False,
    keep_together=False,
):
    paragraph.alignment = alignment
    fmt = paragraph.paragraph_format
    fmt.left_indent = Pt(left) if left is not None else None
    fmt.first_line_indent = Pt(first) if first is not None else None
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.keep_with_next = keep_next
    fmt.keep_together = keep_together
    return paragraph


def clear_cell(cell):
    tc = cell._tc
    for child in list(tc):
        if child.tag != qn("w:tcPr"):
            tc.remove(child)


def clear_direct_paragraphs(cell):
    for child in list(cell._tc):
        if child.tag == qn("w:p"):
            cell._tc.remove(child)


def ensure_paragraph_style(document, name, base="Normal"):
    if name in document.styles:
        style = document.styles[name]
    else:
        style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        if base in document.styles:
            style.base_style = document.styles[base]
    style.font.name = "Times New Roman"
    style.font.size = Pt(13)
    style.font.color.rgb = RGBColor(0, 0, 0)
    return style


def apply_numbering(paragraph, num_id, level=0):
    """Attach a real Word numbering definition rather than a typed bullet/number."""
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = num_pr.find(qn("w:ilvl"))
    if ilvl is None:
        ilvl = OxmlElement("w:ilvl")
        num_pr.append(ilvl)
    ilvl.set(qn("w:val"), str(level))
    num_id_node = num_pr.find(qn("w:numId"))
    if num_id_node is None:
        num_id_node = OxmlElement("w:numId")
        num_pr.append(num_id_node)
    num_id_node.set(qn("w:val"), str(num_id))


def create_numbering_instance(document, abstract_num_id, start=1):
    """Create an independent list instance so numbering always restarts at 1."""
    numbering = document.part.numbering_part.element
    ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    new_id = max(ids, default=0) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(new_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), str(abstract_num_id))
    num.append(abstract)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), str(start))
    override.append(start_override)
    num.append(override)
    numbering.append(num)
    return new_id


def add_paragraph(
    cell,
    text="",
    *,
    style=None,
    bold=False,
    italic=False,
    size=13,
    alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
    left=0,
    first=None,
    before=0,
    after=2,
    line=1.3,
    keep_next=False,
    keep_together=False,
):
    paragraph = cell.add_paragraph(style=style) if style else cell.add_paragraph()
    format_paragraph(
        paragraph,
        alignment=alignment,
        left=left,
        first=first,
        before=before,
        after=after,
        line=line,
        keep_next=keep_next,
        keep_together=keep_together,
    )
    if text:
        run = paragraph.add_run(text)
        set_run_font(run, size=size, bold=bold, italic=italic)
    return paragraph


def add_labelled_paragraph(
    cell,
    label,
    text,
    *,
    alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
    left=0,
    before=0,
    after=2,
    line=1.3,
    size=13,
):
    paragraph = add_paragraph(
        cell,
        alignment=alignment,
        left=left,
        before=before,
        after=after,
        line=line,
    )
    label_run = paragraph.add_run(label)
    set_run_font(label_run, size=size, bold=True)
    body_run = paragraph.add_run(text)
    set_run_font(body_run, size=size)
    return paragraph


def add_heading(cell, text, level=1):
    style = "Heading 1" if level == 1 else "Heading 2"
    paragraph = add_paragraph(
        cell,
        text,
        style=style,
        bold=True,
        size=13,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
        left=12 if level == 1 else 28,
        before=7 if level == 1 else 4,
        after=3,
        line=1.2,
        keep_next=True,
    )
    return paragraph


def add_body(cell, text, *, left=28, bold_prefix=None, italic=False):
    if bold_prefix and text.startswith(bold_prefix):
        return add_labelled_paragraph(
            cell,
            bold_prefix,
            text[len(bold_prefix) :],
            left=left,
        )
    return add_paragraph(cell, text, left=left, italic=italic)


def add_bullet(cell, text, *, level=1, bold_prefix=None):
    style = "List Bullet" if level == 1 else "List Bullet 2"
    left = 46 if level == 1 else 70
    paragraph = add_paragraph(
        cell,
        style=style,
        alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
        left=left,
        before=0,
        after=1.5,
        line=1.25,
        keep_together=True,
    )
    apply_numbering(paragraph, num_id=8, level=level - 1)
    if bold_prefix and text.startswith(bold_prefix):
        run = paragraph.add_run(bold_prefix)
        set_run_font(run, bold=True)
        run = paragraph.add_run(text[len(bold_prefix) :])
        set_run_font(run)
    else:
        run = paragraph.add_run(text)
        set_run_font(run)
    return paragraph


def add_numbered(cell, text):
    paragraph = add_paragraph(
        cell,
        text,
        style="List Number",
        alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
        left=46,
        before=0,
        after=1.5,
        line=1.25,
        keep_together=True,
    )
    apply_numbering(paragraph, num_id=NUMBERED_NUM_ID, level=0)
    return paragraph


def add_figure(cell, filename, caption, width):
    paragraph = add_paragraph(
        cell,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        left=0,
        before=5,
        after=1,
        line=1.0,
        keep_next=True,
        keep_together=True,
    )
    run = paragraph.add_run()
    set_run_font(run, size=11)
    inline_shape = run.add_picture(str(DIAGRAMS / filename), width=Inches(width))
    inline_shape._inline.docPr.set("descr", caption)
    inline_shape._inline.docPr.set("title", caption)
    add_paragraph(
        cell,
        caption,
        italic=True,
        size=11,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        left=0,
        before=0,
        after=5,
        line=1.0,
        keep_together=True,
    )


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def remove_fixed_height(row):
    row.height = None
    row.height_rule = WD_ROW_HEIGHT_RULE.AUTO
    tr_pr = row._tr.get_or_add_trPr()
    for element in list(tr_pr.findall(qn("w:trHeight"))):
        tr_pr.remove(element)


def set_table_cell_content(cell, paragraphs, *, center=False, size=11):
    clear_cell(cell)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for index, item in enumerate(paragraphs):
        if isinstance(item, tuple):
            label, body = item
            paragraph = add_paragraph(
                cell,
                alignment=WD_ALIGN_PARAGRAPH.LEFT,
                left=0,
                before=0,
                after=1.5,
                line=1.1,
                keep_together=True,
            )
            run = paragraph.add_run(label)
            set_run_font(run, size=size, bold=True)
            run = paragraph.add_run(body)
            set_run_font(run, size=size)
        else:
            add_paragraph(
                cell,
                item,
                size=size,
                alignment=WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT,
                left=0,
                before=0,
                after=1.5,
                line=1.1,
                keep_together=True,
            )


def insert_before_table(cell, table, paragraph):
    table._tbl.addprevious(paragraph._p)


if sha256(REFERENCE) != EXPECTED_SHA256:
    raise RuntimeError("Reference DOCX changed; rerun template distillation before authoring.")

for diagram in ("architecture.png", "bpmn_main.png", "job_state.png", "application_state.png", "erd.png"):
    if not (DIAGRAMS / diagram).exists():
        raise FileNotFoundError(DIAGRAMS / diagram)

shutil.copy2(REFERENCE, OUTPUT)
document = Document(OUTPUT)
for style_name in ("List Bullet", "List Bullet 2", "List Number"):
    ensure_paragraph_style(document, style_name)
NUMBERED_NUM_ID = create_numbering_instance(document, abstract_num_id=3, start=1)

if len(document.sections) != 1 or len(document.tables) < 2:
    raise RuntimeError("Reference structure no longer matches artifact contract.")

# Keep the institutional heading and page system. Rewrite only the title/metadata/body slots.
main_table = document.tables[1]

# Project titles.
title_cell = main_table.rows[0].cells[0]
clear_cell(title_cell)
add_labelled_paragraph(
    title_cell,
    "TÊN TIẾNG VIỆT: ",
    "Xây dựng hệ thống UIT Career Hub hỗ trợ kết nối việc làm và quản lý thực tập theo quy trình kiểm duyệt đa vai trò.",
    left=0,
    before=5,
    after=4,
    line=1.2,
)
add_labelled_paragraph(
    title_cell,
    "TÊN TIẾNG ANH: ",
    "Developing UIT Career Hub: A Multi-Role, Approval-Driven Employment and Internship Management Platform.",
    left=0,
    before=2,
    after=5,
    line=1.2,
)

# Advisor, period and team are retained from the supplied template, but normalized to black text.
advisor_cell = main_table.rows[1].cells[0]
clear_cell(advisor_cell)
add_labelled_paragraph(advisor_cell, "Cán bộ hướng dẫn: ", "ThS. Trần Anh Dũng", alignment=WD_ALIGN_PARAGRAPH.LEFT, left=0, before=4, after=4, line=1.15)

period_cell = main_table.rows[2].cells[0]
clear_cell(period_cell)
add_labelled_paragraph(period_cell, "Thời gian thực hiện: ", "Từ 26/01/2026 đến 30/05/2026", alignment=WD_ALIGN_PARAGRAPH.LEFT, left=0, before=4, after=4, line=1.15)

student_cell = main_table.rows[3].cells[0]
clear_cell(student_cell)
add_paragraph(student_cell, "Sinh viên thực hiện:", bold=True, alignment=WD_ALIGN_PARAGRAPH.LEFT, left=0, before=4, after=2, line=1.15)
add_paragraph(student_cell, "Phạm Thị Kiều Diễm - 23520286", alignment=WD_ALIGN_PARAGRAPH.LEFT, left=0, before=0, after=1, line=1.15)
add_paragraph(student_cell, "Nguyễn Trương Ngọc Hân - 23520434", alignment=WD_ALIGN_PARAGRAPH.LEFT, left=0, before=0, after=4, line=1.15)

# Long-form proposal body.
body = main_table.rows[4].cells[0]
clear_cell(body)
add_paragraph(body, "Nội dung đề tài:", bold=True, alignment=WD_ALIGN_PARAGRAPH.LEFT, left=0, before=1, after=3, line=1.15, keep_next=True)

add_heading(body, "1. Lý do chọn đề tài")
add_body(
    body,
    "Trong hoạt động kết nối việc làm và quản lý thực tập tại trường đại học, dữ liệu thường nằm rải rác giữa email, biểu mẫu, bảng tính và các kênh trao đổi riêng. Sinh viên khó theo dõi đầy đủ tình trạng hồ sơ; doanh nghiệp mất thời gian tiếp nhận ứng viên từ nhiều nguồn; bộ phận phụ trách UIT phải kiểm tra thủ công doanh nghiệp, tin tuyển dụng, tài liệu và kết quả thực tập. Khi số lượng hồ sơ tăng, cách làm này dễ tạo ra chậm trễ, thiếu nhất quán và khó truy vết trách nhiệm.",
)
add_body(body, "Các vấn đề thực tiễn chính gồm:", left=28)
add_bullet(body, "Sinh viên thiếu một nơi thống nhất để tìm cơ hội đã được trường kiểm duyệt, quản lý CV và theo dõi toàn bộ tiến trình của từng đơn.")
add_bullet(body, "Doanh nghiệp khó quản lý vòng đời tin tuyển dụng, danh sách ứng viên, lịch phỏng vấn, kết quả và offer trong cùng một hệ thống.")
add_bullet(body, "UIT phải đóng vai trò trung gian kiểm soát chất lượng nhưng thiếu hàng đợi, SLA, lịch sử quyết định và báo cáo tập trung.")
add_bullet(body, "CV, bảng điểm và offer là dữ liệu nhạy cảm; chia sẻ tệp bằng liên kết công khai hoặc gửi qua nhiều kênh làm tăng nguy cơ lộ thông tin.")
add_bullet(body, "Các chuyển trạng thái phức tạp như bổ sung hồ sơ, rút đơn, hủy phỏng vấn, nhận offer và xác nhận nơi thực tập cần được kiểm soát bằng quy tắc rõ ràng để tránh dữ liệu mâu thuẫn.")
add_body(
    body,
    "Từ các vấn đề trên, nhóm lựa chọn xây dựng UIT Career Hub - cổng kết nối việc làm và quản lý thực tập dành cho một trường UIT, với ba vai trò Sinh viên, Bộ phận phụ trách UIT và Doanh nghiệp đối tác. Điểm trọng tâm của đề tài không nằm ở việc tạo thêm một trang đăng tin, mà ở quy trình có kiểm soát: doanh nghiệp được UIT quản lý, tin được duyệt trước khi công khai, hồ sơ được UIT kiểm tra trước khi chuyển doanh nghiệp, offer được đối chiếu trước khi xác nhận placement và mọi quyết định quan trọng đều có lịch sử cùng nhật ký kiểm toán.",
)
add_body(
    body,
    "Giá trị khác biệt của đề tài: ",
    bold_prefix="Giá trị khác biệt của đề tài: ",
)
add_bullet(body, "Thiết kế theo state machine cho cả tin tuyển dụng, đơn ứng tuyển và kỳ thực tập, giúp loại bỏ thao tác cập nhật trạng thái tùy ý.")
add_bullet(body, "Hỗ trợ một sinh viên ứng tuyển nhiều vị trí nhưng chỉ chốt một nơi thực tập; khi UIT xác nhận, các đơn còn hoạt động được đóng nhất quán trong cùng transaction.")
add_bullet(body, "Lưu bản chụp tài liệu theo từng lần nộp để yêu cầu bổ sung không làm mất dấu hồ sơ cũ.")
add_bullet(body, "Áp dụng RBAC, ownership, idempotency, optimistic locking, audit log và URL ký ngắn hạn để bảo vệ dữ liệu và chống thao tác lặp.")

add_heading(body, "2. Mục tiêu")
add_heading(body, "2.1. Mục tiêu tổng quát", level=2)
add_body(
    body,
    "Xây dựng một nền tảng web tập trung hỗ trợ UIT quản lý doanh nghiệp đối tác, kiểm duyệt cơ hội, theo dõi hồ sơ tuyển dụng và kỳ thực tập; đồng thời cung cấp cho sinh viên và doanh nghiệp một quy trình minh bạch, an toàn và có thể truy vết từ lúc đăng tin đến khi hoàn thành đánh giá thực tập.",
)
add_heading(body, "2.2. Mục tiêu cụ thể", level=2)
for item in [
    "Xây dựng cơ chế xác thực và phân quyền cho ba vai trò; kiểm tra quyền ở backend và ràng buộc dữ liệu theo đúng sinh viên hoặc doanh nghiệp sở hữu.",
    "Xây dựng luồng quản lý doanh nghiệp đối tác, tài khoản recruiter và liên kết kích hoạt một lần do UIT cấp.",
    "Xây dựng vòng đời tin tuyển dụng từ nháp, gửi duyệt, yêu cầu chỉnh sửa, duyệt, tạm dừng, hết hạn đến đóng tin.",
    "Cho phép sinh viên quản lý hồ sơ, tải lên tài liệu PDF, chọn CV đã xác minh và nộp đơn theo quy trình hai bước có xác nhận chia sẻ dữ liệu.",
    "Xây dựng luồng UIT kiểm tra hồ sơ, yêu cầu bổ sung, từ chối hoặc chuyển đúng hồ sơ đến đúng doanh nghiệp.",
    "Hỗ trợ doanh nghiệp sàng lọc, chọn Không phù hợp, tạo lịch phỏng vấn, ghi nhận PASS/FAIL và tải PDF offer riêng tư.",
    "Hỗ trợ sinh viên xác nhận lịch, hủy tham gia đúng bước, xem offer, nhận hoặc từ chối; hỗ trợ UIT xác nhận nơi thực tập và đóng các đơn còn lại an toàn.",
    "Theo dõi kỳ thực tập qua HIRED - STARTED - COMPLETED và thu nhận phiếu đánh giá bất biến từ doanh nghiệp và sinh viên.",
    "Cung cấp dashboard theo vai trò, hộp thông báo, quản trị danh mục ngành nghề/kỹ năng và báo cáo có bộ lọc cùng xuất CSV/XLSX.",
    "Xây dựng kiểm thử unit, integration và E2E cho happy flow, nhánh ngoại lệ, RBAC, ownership và các bất biến quan trọng.",
]:
    add_bullet(body, item)

add_heading(body, "2.3. Kết quả mong đợi", level=2)
for item in [
    "Một prototype web có thể trình diễn xuyên suốt ba portal trên cùng dữ liệu PostgreSQL.",
    "Các trạng thái và hành động hiển thị trên giao diện khớp với quy tắc nghiệp vụ ở backend.",
    "Tệp nhạy cảm không công khai; mỗi lần tải xuống được cấp URL ngắn hạn sau khi kiểm tra quyền.",
    "Mọi chuyển trạng thái quan trọng có actor, thời gian, lý do, history, audit và notification phù hợp.",
    "Báo cáo, tài liệu kiến trúc, OpenAPI, state machine, ERD và kịch bản demo đồng bộ với mã nguồn.",
]:
    add_bullet(body, item)

add_heading(body, "3. Phạm vi")
add_heading(body, "3.1. Phạm vi tổ chức và nền tảng", level=2)
add_bullet(body, "Phạm vi một trường UIT; chưa thiết kế multi-school hoặc multi-tenant ở giai đoạn hiện tại.")
add_bullet(body, "Ứng dụng web responsive dùng chung cho Sinh viên, UIT và Doanh nghiệp; chưa phát triển mobile app native.")
add_bullet(body, "Dữ liệu nghiệp vụ tập trung trong PostgreSQL; backend theo kiến trúc modular monolith để bảo đảm transaction xuyên module.")
add_bullet(body, "Tài liệu sinh viên và PDF offer lưu trên Cloudflare R2 private, không lưu trực tiếp trong database.")

add_heading(body, "3.2. Phạm vi chức năng đã bao phủ", level=2)
add_labelled_paragraph(body, "Đối với Sinh viên: ", "đăng nhập bằng tài khoản hợp lệ; xem dashboard; tìm việc và danh bạ đối tác; quản lý hồ sơ, CV và tài liệu; ứng tuyển; theo dõi trạng thái; bổ sung hồ sơ; rút đơn hoặc hủy phỏng vấn đúng bước; phản hồi offer; xem lịch; nhận thông báo; theo dõi placement và gửi phản hồi sau thực tập.", left=28)
add_labelled_paragraph(body, "Đối với UIT: ", "quản lý doanh nghiệp và recruiter; quản trị danh mục; duyệt tin; xác minh tài liệu; duyệt hồ sơ; đối chiếu offer; xác nhận placement; theo dõi HIRED/STARTED/COMPLETED; xem hai phía đánh giá; dashboard; báo cáo và xuất dữ liệu; nhận thông báo xử lý.", left=28)
add_labelled_paragraph(body, "Đối với Doanh nghiệp: ", "kích hoạt tài khoản; cập nhật hồ sơ công khai; quản lý tin; xem ứng viên đúng phạm vi; tải bản chụp hồ sơ được UIT chuyển; sàng lọc; đặt lịch; cập nhật kết quả; quản lý PDF offer; theo dõi dashboard, thông báo và gửi đánh giá thực tập.", left=28)
add_labelled_paragraph(body, "Đối với Hệ thống: ", "JWT/refresh rotation, RBAC/ownership, state history, optimistic lock, idempotency key, notification dedupe, private file access, email outbox, health check, migration, audit và báo cáo.", left=28)

add_heading(body, "3.3. Giới hạn hiện tại và ngoài phạm vi", level=2)
for item in [
    "Chưa tích hợp UIT SSO hoặc đồng bộ trạng thái sinh viên từ hệ thống trường do chưa có contract tích hợp chính thức.",
    "FCM, chat realtime, saved jobs, lịch sử xem tin và nhắc hạn nâng cao chưa thuộc luồng cốt lõi hiện tại.",
    "AI gợi ý việc, AI chấm CV hoặc sinh nội dung không phải chức năng của MVP; chỉ xem xét sau khi quy trình và dữ liệu ổn định.",
    "Báo cáo hiện ưu tiên CSV/XLSX; PDF là phần mở rộng.",
    "Backend chính trong mã nguồn đã chuyển sang Java/Spring Boot; cấu hình Docker/Railway đã có, nhưng endpoint production cần được chốt đồng bộ trước lần phát hành Java chính thức.",
    "Mã Express cũ chỉ được giữ trong thư mục legacy để đối chiếu và không phải kiến trúc mục tiêu của đề tài.",
]:
    add_bullet(body, item)

add_heading(body, "4. Đối tượng sử dụng")
add_heading(body, "4.1. Sinh viên UIT", level=2)
add_body(body, "Sinh viên đang tìm cơ hội việc làm hoặc thực tập, cần quản lý hồ sơ cá nhân, sử dụng tài liệu đã được UIT xác minh và theo dõi minh bạch từng bước của quy trình.")
add_bullet(body, "Chỉ truy cập hồ sơ, tài liệu, đơn, lịch phỏng vấn và đánh giá thuộc chính mình.")
add_bullet(body, "Nhận chỉ dẫn hành động theo trạng thái thực tế thay vì tự chọn trạng thái.")
add_bullet(body, "Có thể ứng tuyển nhiều vị trí nhưng chỉ xác nhận một offer tại một thời điểm.")

add_heading(body, "4.2. Bộ phận phụ trách UIT", level=2)
add_body(body, "Đơn vị vận hành hệ thống và kiểm soát chất lượng đối tác, cơ hội, tài liệu, hồ sơ và kết quả thực tập. Đây là vai trò có phạm vi rộng nhất nhưng mọi mutation quan trọng vẫn phải được ghi audit.")
add_bullet(body, "Quản lý dữ liệu dùng chung và các hàng đợi cần xử lý.")
add_bullet(body, "Đối chiếu thông tin ở các mốc có rủi ro cao: duyệt tin, chuyển hồ sơ, xác minh tài liệu và xác nhận placement.")
add_bullet(body, "Theo dõi số liệu tổng hợp và xuất báo cáo phục vụ quản lý.")

add_heading(body, "4.3. Doanh nghiệp đối tác", level=2)
add_body(body, "Recruiter thuộc doanh nghiệp đã được UIT tạo và kích hoạt, chỉ quản lý tin, ứng viên, offer và đánh giá thuộc đúng doanh nghiệp của mình.")
add_bullet(body, "Không tự đăng ký công khai trong MVP; tài khoản do UIT khởi tạo.")
add_bullet(body, "Không nhìn thấy hồ sơ trước khi UIT chuyển và không thể truy cập chéo doanh nghiệp.")
add_bullet(body, "Chịu trách nhiệm cập nhật kết quả tuyển dụng và đánh giá thực tập đúng thời điểm.")

add_heading(body, "5. Phân tích nghiệp vụ và thiết kế hệ thống")
add_heading(body, "5.1. Luồng nghiệp vụ tổng thể", level=2)
add_body(
    body,
    "Luồng chính bắt đầu từ việc UIT quản lý đối tác, doanh nghiệp gửi tin, UIT duyệt, sinh viên nộp hồ sơ, UIT kiểm tra, doanh nghiệp tuyển chọn, sinh viên phản hồi offer và UIT xác nhận placement. Các nhánh yêu cầu chỉnh sửa, bổ sung, từ chối, rút đơn và hủy phỏng vấn được xử lý bằng endpoint riêng để bảo toàn ý nghĩa nghiệp vụ.",
)
add_figure(body, "bpmn_main.png", "Hình 1. BPMN luồng tuyển dụng và xác nhận thực tập tổng thể", 6.05)

add_heading(body, "5.2. Vòng đời tin tuyển dụng", level=2)
add_body(body, "Tin tuyển dụng không nhận trường status tùy ý từ client. Mỗi hành động có điều kiện, actor và endpoint riêng; chỉ tin RECRUITING còn hạn mới hiển thị cho sinh viên và nhận đơn mới.")
add_figure(body, "job_state.png", "Hình 2. State machine của tin tuyển dụng", 6.05)
add_bullet(body, "Doanh nghiệp tạo DRAFT, gửi PENDING_UIT_REVIEW và xử lý phản hồi REVISION_REQUIRED.")
add_bullet(body, "UIT có thể duyệt thành RECRUITING, yêu cầu sửa hoặc từ chối kèm lý do.")
add_bullet(body, "RECRUITING có thể tạm dừng, hết hạn hoặc đóng; dừng tuyển không xóa các đơn đã nộp.")
add_bullet(body, "Mỗi transition tăng version và ghi lịch sử để chống ghi đè đồng thời.")

add_heading(body, "5.3. Vòng đời đơn ứng tuyển", level=2)
add_body(body, "Đơn ứng tuyển là luồng trung tâm kết nối ba vai trò. Backend khóa bản ghi và kiểm tra trạng thái hiện tại trước mọi quyết định; history và idempotency key giúp retry an toàn mà không tạo side effect trùng.")
add_figure(body, "application_state.png", "Hình 3. State machine của đơn ứng tuyển", 5.45)
add_bullet(body, "UIT_REVIEWING có thể chuyển sang NEEDS_SUPPLEMENT, FORWARDED_TO_COMPANY hoặc UIT_REJECTED.")
add_bullet(body, "Doanh nghiệp chỉ bắt đầu xử lý sau FORWARDED_TO_COMPANY; Không phù hợp và Mời phỏng vấn là hai quyết định riêng.")
add_bullet(body, "PASS tạo bước OFFER_PENDING_STUDENT; sinh viên nhận offer trước, UIT xác nhận placement sau.")
add_bullet(body, "Khi UIT xác nhận HIRED, hệ thống khóa sinh viên và các đơn liên quan, đóng các đơn còn hoạt động với lý do ACCEPTED_OTHER_JOB, hủy lịch liên quan và tạo placement trong cùng transaction.")

add_heading(body, "5.4. Vòng đời kỳ thực tập", level=2)
add_numbered(body, "UIT xác nhận nơi nhận việc và hệ thống tạo placement ở trạng thái HIRED.")
add_numbered(body, "UIT ghi nhận ngày bắt đầu thực tế để chuyển HIRED sang STARTED.")
add_numbered(body, "UIT xác nhận ngày hoàn thành để chuyển STARTED sang COMPLETED.")
add_numbered(body, "Sau COMPLETED, doanh nghiệp và sinh viên mỗi bên gửi tối đa một phiếu bất biến; UIT xem đủ hai phía, còn doanh nghiệp không xem nội dung phản hồi riêng tư của sinh viên.")

add_heading(body, "5.5. Kiến trúc tổng thể", level=2)
add_body(body, "Hệ thống dùng modular monolith để giữ ranh giới module nhưng vẫn cho phép transaction nhất quán ở luồng placement nhiều đơn. Frontend chia portal theo vai trò; backend là điểm quyết định quyền; PostgreSQL lưu dữ liệu có cấu trúc; R2 lưu tài liệu private; email được tách qua outbox để lỗi nhà cung cấp không rollback nghiệp vụ.")
add_figure(body, "architecture.png", "Hình 4. Kiến trúc tổng thể UIT Career Hub", 5.75)

add_heading(body, "5.6. Mô hình dữ liệu mức khái quát", level=2)
add_body(body, "Mô hình dữ liệu xoay quanh users, hồ sơ sinh viên, doanh nghiệp, tin, đơn, tài liệu, phỏng vấn, kết quả và placement. Các bảng history, audit và notification lưu thông tin truy vết, không thay thế dữ liệu nghiệp vụ gốc.")
add_figure(body, "erd.png", "Hình 5. Các thực thể và quan hệ nghiệp vụ chính", 5.65)
add_bullet(body, "Application documents lưu snapshot để hồ sơ đã nộp không thay đổi khi sinh viên cập nhật CV sau đó.")
add_bullet(body, "Unique index ngăn nộp trùng một tin khi đơn cũ còn hoạt động và ngăn một sinh viên đồng thời có hai placement được chấp nhận.")
add_bullet(body, "Placement tách khỏi application để HIRED vẫn là terminal của tuyển dụng trong khi kỳ thực tập tiếp tục STARTED/COMPLETED.")

add_heading(body, "5.7. Yêu cầu phi chức năng", level=2)
add_labelled_paragraph(body, "Bảo mật: ", "JWT có issuer/audience/expiry; refresh token HttpOnly được xoay; mật khẩu băm; RBAC và ownership ở backend; tài nguyên chéo trả 404; response nghiệp vụ không cache.", left=28)
add_labelled_paragraph(body, "An toàn tài liệu: ", "chỉ nhận PDF tối đa 10 MB; kiểm tra MIME, kích thước và chữ ký tệp; object key do server sinh; bucket private; URL ký có thời hạn ngắn và không được ghi log.", left=28)
add_labelled_paragraph(body, "Nhất quán dữ liệu: ", "transaction, SELECT FOR UPDATE, optimistic version, idempotency key, unique index và notification dedupe bảo vệ các thao tác đồng thời.", left=28)
add_labelled_paragraph(body, "Khả dụng: ", "giao diện có loading, empty, error, retry; hỗ trợ desktop/tablet/mobile; dashboard cung cấp công việc cần xử lý theo vai trò.", left=28)
add_labelled_paragraph(body, "Khả năng bảo trì: ", "monorepo tách frontend/backend/database/docs; migration tuần tự; OpenAPI; test tự động; module backend theo domain.", left=28)
add_labelled_paragraph(body, "Truy vết và vận hành: ", "trace ID, structured log, health check, audit log, backup/restore drill và production readiness gate hỗ trợ xử lý sự cố.", left=28)

add_heading(body, "6. Phương pháp thực hiện")
add_heading(body, "6.1. Phương pháp nghiên cứu và phát triển", level=2)
add_bullet(body, "Khảo sát quy trình hiện tại và xác định các điểm bàn giao giữa Sinh viên, UIT và Doanh nghiệp.")
add_bullet(body, "Mô hình hóa domain bằng use case, BPMN, state machine, ERD và ma trận RBAC trước khi cài đặt.")
add_bullet(body, "Phát triển lặp theo lát cắt end-to-end; mỗi lát cắt gồm migration, backend, OpenAPI, UI, audit/notification và test.")
add_bullet(body, "Ưu tiên happy flow và các bất biến khó trước; các tiện ích mở rộng chỉ triển khai sau khi luồng cốt lõi ổn định.")
add_bullet(body, "Đánh giá bằng kiểm thử chức năng, phân quyền, ownership, concurrency và khả năng trình diễn xuyên suốt ba vai trò.")

add_heading(body, "6.2. Quy trình phát triển", level=2)
NUMBERED_NUM_ID = create_numbering_instance(document, abstract_num_id=3, start=1)
for step in [
    "Phân tích yêu cầu, xác định actor, dữ liệu nhạy cảm, trạng thái và điều kiện chuyển trạng thái.",
    "Thiết kế kiến trúc modular monolith, contract API, database schema và chiến lược lưu tệp private.",
    "Cài đặt từng module; kiểm tra validation, RBAC, ownership và transaction ngay tại backend.",
    "Tích hợp UI theo availableActions từ API; bổ sung loading, empty, error, success và confirmation.",
    "Viết unit/integration test và E2E cho happy flow, nhánh lỗi, retry và truy cập chéo.",
    "Chạy typecheck, OpenAPI validation, test, build, migration và kiểm tra demo trước mỗi mốc bàn giao.",
    "Cập nhật báo cáo, state machine, ERD, kịch bản demo và tài liệu vận hành cùng mã nguồn.",
]:
    add_numbered(body, step)

add_heading(body, "6.3. Chiến lược kiểm thử", level=2)
add_bullet(body, "Unit test cho validation, token, service và xử lý lỗi độc lập.")
add_bullet(body, "Integration test với PostgreSQL cho state transition, transaction, unique constraint, audit, notification và idempotency.")
add_bullet(body, "RBAC/ownership negative test cho toàn bộ endpoint giới hạn vai trò và truy cập chéo doanh nghiệp/sinh viên.")
add_bullet(body, "Playwright smoke cho đăng nhập ba vai trò; full E2E cho happy flow, nhánh bổ sung, từ chối, rút đơn, hủy phỏng vấn, từ chối offer, taxonomy và báo cáo.")
add_bullet(body, "Kiểm tra production readiness, bundle, health database và restore trên môi trường test trước khi demo.")

add_heading(body, "7. Nền tảng công nghệ và hiện trạng ứng dụng")
add_heading(body, "7.1. Công nghệ sử dụng", level=2)
add_labelled_paragraph(body, "Frontend: ", "React 19, React DOM 19, Vite 6.4, JavaScript module, Phosphor Icons, Inter và CSS responsive.", left=28)
add_labelled_paragraph(body, "Backend: ", "Java 21, Spring Boot 3.5, Spring Web, Spring Security, OAuth2 Resource Server, Validation, JDBC và Actuator.", left=28)
add_labelled_paragraph(body, "Dữ liệu: ", "PostgreSQL 16/Neon, migration SQL tuần tự và transaction ở service/repository.", left=28)
add_labelled_paragraph(body, "Lưu trữ tệp: ", "Cloudflare R2 private qua S3-compatible SDK; presigned PUT/GET; kiểm tra PDF sau upload.", left=28)
add_labelled_paragraph(body, "Báo cáo: ", "CSV UTF-8 an toàn với Excel và XLSX tạo bằng Apache POI, giới hạn 10.000 dòng.", left=28)
add_labelled_paragraph(body, "Triển khai: ", "frontend cấu hình Vercel; backend Java đóng gói Docker image và có cấu hình Railway/health check; database Neon; R2 cho tệp.", left=28)
add_labelled_paragraph(body, "Kiểm thử và CI: ", "JUnit/Spring Test/Testcontainers, Node test, Playwright, Maven Wrapper, pnpm và GitHub Actions.", left=28)

add_heading(body, "7.2. Hiện trạng chức năng tại thời điểm lập tài liệu", level=2)
for item in [
    "Đã có ba portal và dashboard đọc dữ liệu thật theo phạm vi vai trò.",
    "Đã hoàn thành luồng quản lý đối tác, recruiter, tin tuyển dụng và UIT review queue.",
    "Đã hoàn thành hồ sơ/CV, xác minh tài liệu, ứng tuyển hai bước, bổ sung hồ sơ và snapshot tài liệu.",
    "Đã hoàn thành hàng đợi ứng viên, lịch phỏng vấn, PASS/FAIL, PDF offer private và phản hồi offer.",
    "Đã hoàn thành confirm placement nhiều đơn, vòng đời HIRED/STARTED/COMPLETED và phiếu hai phía.",
    "Đã có notification inbox, daily pending/email outbox, taxonomy admin và báo cáo CSV/XLSX.",
    "Đã có OpenAPI, migration, state machine, RBAC matrix, threat model, kịch bản demo và các bộ kiểm thử tự động.",
]:
    add_bullet(body, item)

add_heading(body, "7.3. Tiêu chí hoàn thành đề tài", level=2)
add_bullet(body, "Happy flow ba vai trò chạy xuyên suốt trên một bộ dữ liệu demo có thể reset ở môi trường riêng.")
add_bullet(body, "Mọi endpoint nghiệp vụ có validation, role check và ownership phù hợp; không lộ tài liệu hoặc tenant khác.")
add_bullet(body, "State machine trong báo cáo, OpenAPI và code khớp nhau; không còn CTA giao diện vượt quá availableActions.")
add_bullet(body, "Bộ test/build đạt; tài liệu bảo vệ có BPMN, ERD, sequence/state, API, RBAC và bằng chứng kiểm thử.")
add_bullet(body, "Bản phát hành Java chốt đúng API URL, secret production được xoay, demo account bị tắt và health check đạt trước khi sử dụng thật.")

# Schedule and future directions: preserve the source nested schedule table and signature furniture.
plan_cell = main_table.rows[5].cells[0]
nested_plan = plan_cell.tables[0]
clear_direct_paragraphs(plan_cell)

plan_heading = add_paragraph(plan_cell, "Kế hoạch thực hiện:", bold=True, alignment=WD_ALIGN_PARAGRAPH.LEFT, left=0, before=4, after=2, line=1.15, keep_next=True)
insert_before_table(plan_cell, nested_plan, plan_heading)
track_a = add_labelled_paragraph(plan_cell, "Track A - Giao diện và trải nghiệm người dùng: ", "Phạm Thị Kiều Diễm", alignment=WD_ALIGN_PARAGRAPH.LEFT, left=28, before=0, after=1, line=1.15)
insert_before_table(plan_cell, nested_plan, track_a)
track_b = add_labelled_paragraph(plan_cell, "Track B - Backend, dữ liệu và an toàn hệ thống: ", "Nguyễn Trương Ngọc Hân", alignment=WD_ALIGN_PARAGRAPH.LEFT, left=28, before=0, after=4, line=1.15)
insert_before_table(plan_cell, nested_plan, track_b)

headers = ["Giai đoạn", "Thời gian", "Công việc", "Kết quả"]
for index, text in enumerate(headers):
    set_table_cell_content(nested_plan.rows[0].cells[index], [text], center=True, size=11.5)
    for run in nested_plan.rows[0].cells[index].paragraphs[0].runs:
        run.bold = True
set_repeat_table_header(nested_plan.rows[0])

schedule = [
    (
        "Khởi động và chốt phạm vi",
        "26/01 - 09/02/2026",
        [("Track A: ", "Khảo sát hành trình ba vai trò, lập inventory màn hình và prototype luồng chính."), ("Track B: ", "Xác định actor, trạng thái, dữ liệu nhạy cảm, kiến trúc và contract API sơ bộ.")],
        "Phạm vi MVP, backlog, BPMN tổng thể và tiêu chí demo được thống nhất.",
    ),
    (
        "Thiết kế domain và dữ liệu",
        "10/02 - 01/03/2026",
        [("Track A: ", "Thiết kế navigation, form, bảng, modal và trạng thái loading/empty/error."), ("Track B: ", "Thiết kế ERD, migration, state machine Job/Application, RBAC và ownership.")],
        "Có wireframe, schema, API contract và state machine nền tảng.",
    ),
    (
        "Xác thực, đối tác và tin tuyển dụng",
        "02/03 - 23/03/2026",
        [("Track A: ", "Hoàn thiện login và portal UIT/Company cho doanh nghiệp, recruiter, tin và review queue."), ("Track B: ", "Cài đặt JWT/refresh, partner management, job lifecycle, history, audit và test.")],
        "Doanh nghiệp tạo/gửi tin; UIT duyệt, yêu cầu sửa hoặc từ chối an toàn.",
    ),
    (
        "Hồ sơ và ứng tuyển",
        "24/03 - 14/04/2026",
        [("Track A: ", "Hoàn thiện hồ sơ/CV, tìm việc, ứng tuyển hai bước và theo dõi đơn."), ("Track B: ", "Cài đặt R2 private, xác minh tài liệu, snapshot, supplement, withdraw và UIT review.")],
        "Luồng Student - UIT - Company chạy đến bước chuyển hồ sơ.",
    ),
    (
        "Phỏng vấn, offer và placement",
        "15/04 - 05/05/2026",
        [("Track A: ", "Hoàn thiện candidate pipeline, lịch, PDF offer, phản hồi và placement UI."), ("Track B: ", "Cài đặt company decisions, interview, PASS/FAIL, offer, confirm placement và concurrency guard.")],
        "Happy flow hoàn chỉnh đến HIRED; các đơn khác đóng nhất quán.",
    ),
    (
        "Quản trị, báo cáo và hardening",
        "06/05 - 25/05/2026",
        [("Track A: ", "Hoàn thiện dashboard, taxonomy, reports, placement lifecycle, evaluation và responsive."), ("Track B: ", "Cài đặt report/export, notification/email outbox, security hardening, health và backup drill.")],
        "MVP đủ chức năng quản trị, báo cáo, theo dõi thực tập và vận hành cơ bản.",
    ),
    (
        "Kiểm thử, báo cáo và demo",
        "26/05 - 30/05/2026",
        [("Cả nhóm: ", "Chạy unit/integration/E2E, rà UI, chốt Java deployment, chuẩn bị dữ liệu, báo cáo, slide, video và kịch bản bảo vệ.")],
        "Hệ thống ổn định để demo; tài liệu kỹ thuật và hồ sơ bảo vệ đồng bộ với code.",
    ),
]

for row_index, item in enumerate(schedule, start=1):
    row = nested_plan.rows[row_index]
    remove_fixed_height(row)
    phase, dates, tasks, result = item
    set_table_cell_content(row.cells[0], [phase], center=True, size=11)
    set_table_cell_content(row.cells[1], [dates], center=True, size=11)
    set_table_cell_content(row.cells[2], tasks, center=False, size=10.8)
    set_table_cell_content(row.cells[3], [result], center=False, size=10.8)

add_heading(plan_cell, "8. Hướng phát triển đề tài")
add_heading(plan_cell, "8.1. Hoàn thiện tích hợp và trải nghiệm", level=2)
add_bullet(plan_cell, "Tích hợp UIT SSO và đồng bộ trạng thái sinh viên khi có API, quy định và dữ liệu chính thức từ trường.")
add_bullet(plan_cell, "Bổ sung saved jobs, lịch sử xem tin, nhắc hạn nộp và trải nghiệm tìm kiếm nâng cao.")
add_bullet(plan_cell, "Hoàn thành accessibility audit cho keyboard, focus, screen reader, zoom 200% và contrast; tăng test component/interaction cho frontend.")
add_bullet(plan_cell, "Tách các component lớn theo feature và tiếp tục giữ lazy loading theo portal để duy trì bundle hợp lý.")

add_heading(plan_cell, "8.2. Vận hành và bảo mật production", level=2)
add_bullet(plan_cell, "Chốt deployment backend Java, cập nhật rewrite frontend, kiểm tra health/readiness và phương án rollback.")
add_bullet(plan_cell, "Bổ sung rate limit theo user/IP cho mutation tốn tài nguyên, quy định retention và quyền truy cập structured log.")
add_bullet(plan_cell, "Hoàn thiện theo dõi email delivered/bounced, SchedulerLog, retry dashboard và FCM nếu có nhu cầu thực tế.")
add_bullet(plan_cell, "Tích hợp quét malware/quarantine cho PDF trước khi mở cho người dùng thật; xoay secret và loại tài khoản demo.")

add_heading(plan_cell, "8.3. Mở rộng nghiệp vụ sau MVP", level=2)
add_bullet(plan_cell, "Báo cáo PDF, phân quyền chi tiết theo chức danh UIT và nhiều recruiter với quyền khác nhau.")
add_bullet(plan_cell, "Nghiên cứu AI gợi ý việc hoặc hỗ trợ CV trên dữ liệu đã chuẩn hóa, có đánh giá bias, giải thích và quyền riêng tư; không để AI thay thế quyết định tuyển dụng.")
add_bullet(plan_cell, "Chỉ xem xét mobile app, chat realtime hoặc multi-school khi kiến trúc tenant, bảo mật và nguồn lực vận hành đã được xác định rõ.")

# Normalize the signature block to plain black text while preserving its geometry and blank signing space.
left_signature = main_table.rows[6].cells[0]
clear_cell(left_signature)
add_paragraph(left_signature, "Xác nhận của CBHD", bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER, left=0, before=6, after=4, line=1.15)
add_paragraph(left_signature, "(Ký tên và ghi rõ họ tên)", alignment=WD_ALIGN_PARAGRAPH.CENTER, left=0, before=0, after=6, line=1.15)

right_signature = main_table.rows[6].cells[1]
clear_cell(right_signature)
add_paragraph(right_signature, "TP. HCM, ngày .... tháng .... năm ....", bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER, left=0, before=6, after=3, line=1.15)
add_paragraph(right_signature, "Nhóm Sinh viên", bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER, left=0, before=0, after=2, line=1.15)
add_paragraph(right_signature, "(Ký tên và ghi rõ họ tên)", alignment=WD_ALIGN_PARAGRAPH.CENTER, left=0, before=0, after=6, line=1.15)

# Explicitly keep the source page geometry and black/white system.
section = document.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.header_distance = Inches(0.5)
section.footer_distance = Inches(0.5)

document.core_properties.title = "Đề cương chi tiết - UIT Career Hub"
document.core_properties.subject = "Cổng kết nối việc làm và quản lý thực tập UIT"
document.core_properties.keywords = "UIT Career Hub, tuyển dụng, thực tập, state machine, RBAC"
document.core_properties.author = "Nhóm sinh viên 23520286 - 23520434"

document.save(OUTPUT)
print(f"Created {OUTPUT}")
print(f"Reference SHA-256: {sha256(REFERENCE)}")
print(f"Output size: {OUTPUT.stat().st_size} bytes")
