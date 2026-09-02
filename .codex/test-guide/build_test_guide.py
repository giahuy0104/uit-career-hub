from __future__ import annotations

from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUTPUT = ROOT / "UIT_Career_Hub_Quy_Trinh_Kiem_Thu_3_Vai_Tro.docx"

# compact_reference_guide with a named monochrome academic override.
FONT = "Times New Roman"
BLACK = "000000"
GRAY = "F2F2F2"
MID_GRAY = "D9D9D9"
WHITE = "FFFFFF"
CONTENT_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGIN = {"top": 90, "bottom": 90, "start": 120, "end": 120}


def set_run_font(run, size=11, *, bold=None, italic=None, color=BLACK):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), FONT)


def set_style_font(style, size, *, bold=False, color=BLACK):
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    for theme_key in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        attribute = qn(f"w:{theme_key}")
        if attribute in fonts.attrib:
            del fonts.attrib[attribute]
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), FONT)


def configure_styles(doc: Document):
    styles = doc.styles
    normal = styles["Normal"]
    set_style_font(normal, 11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    title = styles["Title"]
    set_style_font(title, 24, bold=True)
    title.font.italic = False
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(8)
    title.paragraph_format.keep_with_next = True

    subtitle = styles["Subtitle"]
    set_style_font(subtitle, 12)
    subtitle.font.italic = False
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(12)

    heading_tokens = {
        "Heading 1": (16, 18, 10),
        "Heading 2": (13, 14, 7),
        "Heading 3": (11.5, 10, 5),
    }
    for name, (size, before, after) in heading_tokens.items():
        style = styles[name]
        set_style_font(style, size, bold=True)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.05
        style.paragraph_format.keep_with_next = True

    # Remove decorative borders inherited from built-in Word title styles.
    for style_name in ("Title", "Subtitle"):
        ppr = styles[style_name].element.get_or_add_pPr()
        border = ppr.find(qn("w:pBdr"))
        if border is not None:
            ppr.remove(border)

    for name, size, italic in (
        ("Guide Note", 10, False),
        ("Table Text", 9.2, False),
        ("Table Header", 9.2, False),
        ("Small Text", 9.5, False),
    ):
        if name not in styles:
            styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style = styles[name]
        set_style_font(style, size, bold=name == "Table Header")
        style.font.italic = italic
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(2 if "Table" in name else 4)
        style.paragraph_format.line_spacing = 1.12 if "Table" in name else 1.2


def set_cell_shading(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, margins=CELL_MARGIN):
    tcpr = cell._tc.get_or_add_tcPr()
    tcmar = tcpr.first_child_found_in("w:tcMar")
    if tcmar is None:
        tcmar = OxmlElement("w:tcMar")
        tcpr.append(tcmar)
    for side, value in margins.items():
        node = tcmar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tcmar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tcpr = cell._tc.get_or_add_tcPr()
    tcw = tcpr.find(qn("w:tcW"))
    if tcw is None:
        tcw = OxmlElement("w:tcW")
        tcpr.append(tcw)
    tcw.set(qn("w:w"), str(width_dxa))
    tcw.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    trpr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "1")
    trpr.append(tbl_header)


def prevent_row_split(row):
    trpr = row._tr.get_or_add_trPr()
    if trpr.find(qn("w:cantSplit")) is None:
        node = OxmlElement("w:cantSplit")
        node.set(qn("w:val"), "1")
        trpr.append(node)


def set_table_geometry(table, widths: Sequence[int], indent=TABLE_INDENT_DXA):
    assert sum(widths) == CONTENT_DXA, (widths, sum(widths))
    table.autofit = False
    tblpr = table._tbl.tblPr
    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.append(tblw)
    tblw.set(qn("w:w"), str(CONTENT_DXA))
    tblw.set(qn("w:type"), "dxa")
    tblind = tblpr.find(qn("w:tblInd"))
    if tblind is None:
        tblind = OxmlElement("w:tblInd")
        tblpr.append(tblind)
    tblind.set(qn("w:w"), str(indent))
    tblind.set(qn("w:type"), "dxa")
    layout = tblpr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblpr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def table_borders(table, color=BLACK, size=6):
    tblpr = table._tbl.tblPr
    borders = tblpr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblpr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def clear_paragraph(paragraph):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def fill_cell(cell, text, *, header=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = cell.paragraphs[0]
    clear_paragraph(p)
    p.style = "Table Header" if header else "Table Text"
    p.alignment = align
    p.paragraph_format.keep_together = True
    run = p.add_run(str(text))
    set_run_font(run, 9.2, bold=header)
    return p


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    begin_run = OxmlElement("w:r")
    begin_run.append(begin)
    instr_run = OxmlElement("w:r")
    instr_run.append(instr)
    separate_run = OxmlElement("w:r")
    separate_run.append(separate)
    text_run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "1"
    text_run.append(text)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run = OxmlElement("w:r")
    end_run.append(end)
    run = paragraph.add_run("Trang ")
    set_run_font(run, 9)
    paragraph._p.append(begin_run)
    paragraph._p.append(instr_run)
    paragraph._p.append(separate_run)
    paragraph._p.append(text_run)
    paragraph._p.append(end_run)


def set_document_furniture(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("UIT CAREER HUB  |  QUY TRÌNH KIỂM THỬ 3 VAI TRÒ")
    set_run_font(run, 8.5, bold=True)

    footer = section.footer
    add_page_field(footer.paragraphs[0])


def create_numbering(doc: Document, *, fmt: str, marker: str, bullet_font=None):
    root = doc.part.numbering_part.element
    abstract_ids = [int(n.get(qn("w:abstractNumId"))) for n in root.findall(qn("w:abstractNum"))]
    num_ids = [int(n.get(qn("w:numId"))) for n in root.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    numfmt = OxmlElement("w:numFmt")
    numfmt.set(qn("w:val"), fmt)
    lvl.append(numfmt)
    lvltext = OxmlElement("w:lvlText")
    lvltext.set(qn("w:val"), marker)
    lvl.append(lvltext)
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    lvl.append(suff)
    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    ppr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    ppr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    ppr.append(spacing)
    lvl.append(ppr)
    if bullet_font:
        rpr = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:ascii"), bullet_font)
        fonts.set(qn("w:hAnsi"), bullet_font)
        rpr.append(fonts)
        lvl.append(rpr)
    abstract.append(lvl)
    first_num = root.find(qn("w:num"))
    if first_num is None:
        root.append(abstract)
    else:
        first_num.addprevious(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    root.append(num)
    return num_id


def apply_numbering(paragraph, num_id):
    ppr = paragraph._p.get_or_add_pPr()
    numpr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    numpr.append(ilvl)
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    numpr.append(num)
    ppr.append(numpr)


def add_bullet(doc, text, bullet_num):
    p = doc.add_paragraph()
    apply_numbering(p, bullet_num)
    run = p.add_run(text)
    set_run_font(run)
    return p


def add_numbered(doc, text, decimal_num):
    p = doc.add_paragraph()
    apply_numbering(p, decimal_num)
    run = p.add_run(text)
    set_run_font(run)
    return p


def add_labelled(doc, label, text, *, style=None):
    p = doc.add_paragraph(style=style)
    run = p.add_run(label)
    set_run_font(run, bold=True)
    run = p.add_run(text)
    set_run_font(run)
    return p


def add_note(doc, label, text):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_DXA])
    table_borders(table, color=BLACK, size=6)
    cell = table.cell(0, 0)
    set_cell_shading(cell, WHITE)
    p = fill_cell(cell, "", align=WD_ALIGN_PARAGRAPH.LEFT)
    p.style = "Guide Note"
    run = p.add_run(label)
    set_run_font(run, 10, bold=True)
    run = p.add_run(text)
    set_run_font(run, 10)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_simple_table(doc, headers, rows, widths, *, font_size=9.2, center_columns=()):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        fill_cell(table.rows[0].cells[i], header, header=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(table.rows[0].cells[i], GRAY)
    set_repeat_table_header(table.rows[0])
    for row_data in rows:
        row = table.add_row()
        prevent_row_split(row)
        for i, value in enumerate(row_data):
            p = fill_cell(
                row.cells[i],
                value,
                align=WD_ALIGN_PARAGRAPH.CENTER if i in center_columns else WD_ALIGN_PARAGRAPH.LEFT,
            )
            for run in p.runs:
                set_run_font(run, font_size)
    set_table_geometry(table, widths)
    table_borders(table)
    after = doc.add_paragraph()
    after.paragraph_format.space_before = Pt(2)
    after.paragraph_format.space_after = Pt(0)
    return table


def add_test_table(doc, title, cases, *, page_before=False, chunk_size=6):
    chunks = [cases[index:index + chunk_size] for index in range(0, len(cases), chunk_size)]
    for index, chunk in enumerate(chunks):
        if page_before or index > 0:
            doc.add_page_break()
        label = title if index == 0 else f"{title} (tiếp)"
        doc.add_heading(label, level=2)
        add_simple_table(
            doc,
            ["ID", "Kịch bản / điều kiện", "Các bước chính", "Kết quả mong đợi"],
            chunk,
            [1100, 2200, 3000, 3060],
            font_size=9.0,
            center_columns=(0,),
        )
        page_before = False


def add_chunked_matrix(doc, headers, rows, widths, *, chunk_size=6, font_size=9.0, center_columns=(), continuation_title=None):
    chunks = [rows[index:index + chunk_size] for index in range(0, len(rows), chunk_size)]
    for index, chunk in enumerate(chunks):
        if index > 0:
            doc.add_page_break()
            if continuation_title:
                doc.add_heading(f"{continuation_title} (tiếp)", level=2)
        add_simple_table(doc, headers, chunk, widths, font_size=font_size, center_columns=center_columns)


def add_status_legend(doc):
    add_simple_table(
        doc,
        ["Ký hiệu", "Ý nghĩa", "Cách ghi bằng chứng"],
        [
            ("P0", "Luồng chính hoặc ràng buộc an toàn; lỗi làm hệ thống không thể demo.", "Screenshot + response API + trạng thái sau test."),
            ("P1", "Chức năng quan trọng theo vai trò; cần pass trước khi bàn giao.", "Screenshot và ghi chú dữ liệu đã dùng."),
            ("P2", "Khả dụng, hiển thị, empty/error state hoặc tính tiện dụng.", "Screenshot hoặc video ngắn."),
            ("PASS / FAIL / BLOCKED", "Kết quả thực thi thực tế.", "Ghi defect ID nếu FAIL; ghi blocker và người phụ trách nếu BLOCKED."),
        ],
        [1440, 3900, 4020],
        center_columns=(0,),
    )


def add_test_run_record(doc):
    doc.add_heading("Biên bản một lần chạy test", level=2)
    add_simple_table(
        doc,
        ["Trường", "Nội dung ghi nhận"],
        [
            ("Mã lần chạy", "____________________"),
            ("Ngày / người test", "____________________"),
            ("Commit / branch", "____________________"),
            ("Frontend / Backend URL", "____________________"),
            ("Database / checkpoint", "____________________"),
            ("Trình duyệt / viewport", "____________________"),
            ("Tổng PASS / FAIL / BLOCKED", "____________________"),
            ("Defect nghiêm trọng", "____________________"),
            ("Kết luận", "☐ Đủ điều kiện demo    ☐ Cần sửa và chạy lại"),
        ],
        [2700, 6660],
    )


AUTH_CASES = [
    ("AUTH-01", "[P0] Sinh viên đã ACTIVE; email/mật khẩu hợp lệ.", "Đăng nhập; quan sát URL, cookie refresh và dashboard.", "Vào portal Sinh viên; API login/dashboard 200; không hiện menu UIT/Company."),
    ("AUTH-02", "[P0] UIT Admin đã ACTIVE.", "Đăng nhập bằng tài khoản admin.career@uit.edu.vn.", "Vào Tổng quan vận hành; nhìn đủ 9 phân hệ UIT."),
    ("AUTH-03", "[P0] Recruiter thuộc doanh nghiệp ACTIVE.", "Đăng nhập bằng recruiter@vng.example.", "Vào Tổng quan tuyển dụng; dữ liệu chỉ thuộc VNG."),
    ("AUTH-04", "[P1] Sai mật khẩu, email trống/sai định dạng.", "Gửi form với từng dữ liệu sai; thử nhiều lần trong giới hạn.", "Hiện lỗi dễ hiểu; không tạo phiên; không lộ tài khoản có tồn tại hay không."),
    ("AUTH-05", "[P0] Phiên hợp lệ, access token hết hạn.", "Giữ refresh cookie; gọi API để client tự refresh; lặp lại một request.", "Access token được xoay an toàn; request tiếp tục; refresh token cũ không dùng lại được."),
    ("AUTH-06", "[P0] Đang đăng nhập.", "Chọn Đăng xuất; sau đó mở lại màn hình bảo vệ.", "Logout 204; refresh token bị thu hồi; UI trở về đăng nhập; API bảo vệ từ chối."),
    ("AUTH-07", "[P0] User đang có token nhưng bị chuyển SUSPENDED.", "Dùng token cũ gọi /auth/me và một API nghiệp vụ.", "Trả 401 AUTH_ACCESS_REVOKED; không thao tác dữ liệu."),
    ("AUTH-08", "[P0] Recruiter mới có activation token còn hạn.", "Mở link, đặt mật khẩu hợp lệ; gửi lại cùng token.", "Kích hoạt thành công một lần; lần hai/expired token bị từ chối; audit được ghi."),
    ("AUTH-09", "[P0] Có token hợp lệ nhưng sai role hoặc thiếu studentProfileId/companyId.", "Gọi endpoint riêng của role khác và endpoint cần context.", "Sai role trả 403 AUTH_FORBIDDEN; thiếu context trả 403 AUTH_CONTEXT_MISSING."),
]


STUDENT_MODULES = [
    ("6.1. Tổng quan và điều hướng", [
        ("STU-01", "[P1] Sinh viên có hồ sơ và nhiều đơn.", "Mở Tổng quan; đối chiếu % hồ sơ, đơn active, lịch sắp tới, offer chờ.", "Số liệu khớp database/API và thay đổi sau transition."),
        ("STU-02", "[P1] Có task bổ sung, phỏng vấn hoặc offer.", "Bấm từng việc cần làm và các CTA Xem đơn/Xem tất cả.", "Đi đúng màn hình, chọn đúng resourceId; không mất phiên."),
        ("STU-03", "[P2] Tài khoản mới không có đơn/cơ hội.", "Mở dashboard; mô phỏng API chậm và lỗi rồi chọn Thử lại.", "Loading, empty và error state rõ ràng; retry tải lại thành công."),
        ("STU-04", "[P2] Viewport 1440 px, 768 px và 390 px.", "Duyệt toàn bộ menu; thu/phóng; dùng keyboard.", "Menu/CTA không tràn; focus thấy rõ; không có cuộn ngang ngoài ý muốn."),
    ]),
    ("6.2. Việc làm và doanh nghiệp đối tác", [
        ("STU-05", "[P1] Có nhiều tin RECRUITING.", "Tìm theo tên; đổi bộ lọc loại cơ hội/hình thức/địa điểm.", "Danh sách trả đúng tập giao; xóa lọc khôi phục đầy đủ."),
        ("STU-06", "[P1] Chọn một tin đang tuyển.", "Mở chi tiết; đối chiếu title, company, deadline, work mode, mô tả, yêu cầu, quyền lợi.", "Chi tiết khớp tin đã được UIT duyệt; CTA ứng tuyển đúng trạng thái."),
        ("STU-07", "[P0] Có tin DRAFT/PENDING/REJECTED/CLOSED/EXPIRED trong DB.", "Thử truy vấn danh sách công khai và mở trực tiếp ID.", "Chỉ RECRUITING còn hạn được công khai; tin khác không lộ."),
        ("STU-08", "[P1] Sinh viên đã nộp cùng tin hoặc tin hết hạn.", "Mở chi tiết và quan sát CTA; thử gửi request trực tiếp.", "UI không cho nộp lại; backend từ chối mà không tạo đơn trùng."),
        ("STU-09", "[P1] Có ít nhất hai đối tác ACTIVE.", "Mở Doanh nghiệp đối tác; tìm kiếm; mở hồ sơ.", "Hiển thị đúng tên/logo/lĩnh vực/mô tả/website và tin đang tuyển."),
        ("STU-10", "[P0] Có doanh nghiệp/recruiter bị SUSPENDED.", "Tìm doanh nghiệp bị ngưng; kiểm tra dữ liệu công khai.", "Không lộ email recruiter, activation token hoặc dữ liệu nội bộ; đối tác không hợp lệ không xuất hiện."),
    ]),
    ("6.3. Hồ sơ, CV và tài liệu", [
        ("STU-11", "[P1] Hồ sơ sinh viên đã seed.", "Mở Hồ sơ & CV; đối chiếu MSSV, email, khoa, ngành, GPA, điện thoại.", "Dữ liệu đúng studentProfileId của user; trường chỉ đọc không sửa được."),
        ("STU-12", "[P1] Số điện thoại hợp lệ/không hợp lệ.", "Cập nhật số điện thoại; lưu; refresh; thử chuỗi quá dài.", "Dữ liệu hợp lệ được lưu; validation chặn dữ liệu sai; version/audit thay đổi."),
        ("STU-13", "[P0] PDF hợp lệ nhỏ hơn 10 MB.", "Tạo upload, PUT tệp lên R2, complete upload; làm mới danh sách.", "Tài liệu xuất hiện PENDING; metadata đúng; storage key không trả về client."),
        ("STU-14", "[P0] Tệp PNG/EXE, PDF >10 MB, MIME/signature không khớp.", "Thử upload từng tệp và gọi complete.", "Request bị từ chối; không tạo tài liệu hợp lệ; không để object mồ côi kéo dài."),
        ("STU-15", "[P1] Có hai CV đã VERIFIED.", "Chọn CV mặc định; refresh; chọn CV còn lại.", "Chỉ một CV default; thay đổi được lưu nhất quán."),
        ("STU-16", "[P0] Tài liệu thuộc sinh viên.", "Chọn Mở/Tải; ghi thời điểm URL; thử dùng sau khi hết hạn.", "Nhận URL ký ngắn hạn; URL hết hạn bị từ chối; lần cấp URL được audit."),
        ("STU-17", "[P1] Có tài liệu chưa dùng và tài liệu đã snapshot trong đơn.", "Xóa tài liệu gốc; mở lại đơn cũ.", "Danh sách hồ sơ cập nhật; snapshot của đơn cũ không bị thay đổi."),
    ]),
    ("6.4. Đơn ứng tuyển và offer", [
        ("STU-18", "[P0] Tin RECRUITING và CV/tài liệu đủ điều kiện.", "Chọn Ứng tuyển; xem bước chia sẻ dữ liệu; không tick consent rồi thử tiếp tục.", "Không cho gửi nếu chưa đồng ý; hiển thị đúng tài liệu sẽ chia sẻ."),
        ("STU-19", "[P0] Đã hoàn tất bước xác nhận.", "Gửi đơn; quan sát response và danh sách đơn.", "Tạo một đơn UIT_REVIEWING; có snapshot tài liệu, history, audit và notification UIT."),
        ("STU-20", "[P0] Cùng student/job vẫn còn đơn active.", "Double-click Gửi; retry cùng Idempotency-Key; gửi key khác.", "Cùng key trả cùng kết quả; key khác không tạo đơn active trùng."),
        ("STU-21", "[P1] Có nhiều đơn/trạng thái.", "Tìm theo vị trí; lọc trạng thái; mở chi tiết và history.", "Danh sách, nhãn trạng thái, version, actor, reason và thời gian khớp backend."),
        ("STU-22", "[P0] Đơn NEEDS_SUPPLEMENT có yêu cầu từ UIT.", "Mở yêu cầu; bổ sung note/tài liệu; chọn Gửi lại.", "Đơn về UIT_REVIEWING; snapshot mới được tạo; bản nộp cũ vẫn truy vết được."),
        ("STU-23", "[P0] Đơn ở trạng thái cho phép rút.", "Chọn Rút đơn; nhập lý do hợp lệ; xác nhận.", "Đơn thành WITHDRAWN; history/reason/notification đủ; CTA chuyển trạng thái biến mất."),
        ("STU-24", "[P0] Đơn terminal hoặc đã qua mốc không cho rút.", "Gửi withdraw/accept/decline sai trạng thái qua UI và API.", "Backend từ chối transition; version/history không thay đổi."),
        ("STU-25", "[P0] Đơn OFFER_PENDING_STUDENT có PDF offer.", "Mở PDF; chọn Nhận offer; xác nhận modal.", "PDF mở bằng URL ký; đơn thành ACCEPTED_PENDING_UIT_CONFIRMATION; UIT được thông báo."),
        ("STU-26", "[P0] Đơn OFFER_PENDING_STUDENT.", "Chọn Từ chối offer; nhập lý do; xác nhận.", "Đơn thành OFFER_DECLINED; doanh nghiệp nhận thông báo; không tạo placement."),
        ("STU-27", "[P0] Sinh viên có nhiều đơn active.", "Nhận một offer nhưng chưa được UIT xác nhận.", "Các đơn khác vẫn active; chỉ UIT confirm-placement mới đóng đơn khác."),
    ]),
    ("6.5. Lịch phỏng vấn, thông báo và đánh giá", [
        ("STU-28", "[P1] Có lịch sắp tới và lịch sử.", "Mở Lịch phỏng vấn; đối chiếu thời gian, mode, địa điểm/link, interviewer, kết quả.", "Lịch được tách đúng upcoming/history và chỉ thuộc sinh viên."),
        ("STU-29", "[P0] Lịch PENDING_STUDENT_CONFIRMATION.", "Chọn Xác nhận tham gia; refresh.", "Lịch thành CONFIRMED; CTA đúng; doanh nghiệp nhận cập nhật."),
        ("STU-30", "[P0] Lịch chưa hoàn tất.", "Chọn Không thể tham gia; nhập lý do >=5 ký tự; xác nhận.", "Đơn chuyển WITHDRAWN, lịch CANCELLED; lý do/history/notification nhất quán."),
        ("STU-31", "[P1] Có thông báo đã/chưa đọc.", "Mở Thông báo; lọc; đánh dấu một tin và đọc tất cả.", "Unread count giảm chính xác; chỉ cập nhật thông báo của user."),
        ("STU-32", "[P1] Notification có resourceId.", "Bấm mở thông báo về đơn/lịch/offer.", "Đi đúng màn hình và tự chọn đúng bản ghi."),
        ("STU-33", "[P1] Placement đã COMPLETED, Company đã gửi phiếu.", "Mở đơn HIRED; xem timeline placement và kết quả đánh giá doanh nghiệp.", "Hiển thị HIRED/STARTED/COMPLETED và đánh giá Company đúng dữ liệu."),
        ("STU-34", "[P0] Placement COMPLETED và chưa có phiếu Student.", "Chọn điểm 4 tiêu chí; nhập strengths/improvements; gửi hai lần.", "Lần đầu 201; lần hai không tạo phiếu trùng; nội dung Student chỉ hiển thị cho Student và UIT."),
        ("STU-35", "[P0] Có hai sinh viên khác nhau.", "Dùng Student A truy cập ID đơn, tài liệu, lịch và phiếu của Student B.", "Trả 404 tương ứng; không lộ sự tồn tại hoặc metadata."),
    ]),
]


UIT_MODULES = [
    ("7.1. Tổng quan vận hành", [
        ("UIT-01", "[P1] Checkpoint có tin/hồ sơ/placement đang chờ.", "Mở dashboard; đối chiếu 4 metric, overdue 24h, funnel và hires trong tháng.", "Số liệu khớp truy vấn dashboard và thay đổi sau khi xử lý."),
        ("UIT-02", "[P1] Mỗi hàng đợi có dữ liệu.", "Bấm hồ sơ chờ, tin chờ và placement chờ trên dashboard.", "Đi đúng phân hệ; danh sách ưu tiên bản ghi chờ lâu."),
        ("UIT-03", "[P2] Không có hàng đợi hoặc API lỗi.", "Mở dashboard; mô phỏng empty/error; chọn Thử lại.", "Empty copy đúng; error không làm hỏng layout; retry phục hồi."),
    ]),
    ("7.2. Doanh nghiệp đối tác và recruiter", [
        ("UIT-04", "[P1] Có nhiều đối tác/trạng thái.", "Tìm kiếm; lọc ACTIVE/SUSPENDED; mở chi tiết.", "Danh sách, company profile, recruiter và version đúng phạm vi."),
        ("UIT-05", "[P0] Dữ liệu company và recruiter mới hợp lệ.", "Tạo đối tác kèm primary recruiter; lưu activation link.", "Tạo company/recruiter một transaction; token thô chỉ hiển thị theo contract; audit đủ."),
        ("UIT-06", "[P1] Mã đối tác, email hoặc dữ liệu bắt buộc bị trùng/sai.", "Gửi form từng trường hợp; kiểm tra thông báo.", "Validation/chỉ mục unique chặn; không tạo company/recruiter dở dang."),
        ("UIT-07", "[P0] Hai cửa sổ cùng mở một company.", "Cửa sổ A lưu; cửa sổ B lưu expectedVersion cũ.", "A thành công; B nhận conflict; dữ liệu A không bị ghi đè."),
        ("UIT-08", "[P0] Company ACTIVE có recruiter.", "Tạm ngưng company kèm lý do; dùng token recruiter cũ.", "Company/recruiter bị hạn chế theo rule; token cũ bị từ chối; audit/notification được ghi."),
        ("UIT-09", "[P0] Company SUSPENDED.", "Khôi phục company; đăng nhập lại recruiter hợp lệ.", "Company về ACTIVE; quyền hoạt động trở lại theo trạng thái recruiter."),
        ("UIT-10", "[P1] Company ACTIVE.", "Thêm recruiter thứ hai; kiểm tra activation flow.", "Recruiter gắn đúng companyId; chưa kích hoạt không đăng nhập được."),
        ("UIT-11", "[P0] Recruiter ACTIVE.", "Tạm ngưng/khôi phục recruiter; kiểm tra token và dữ liệu company.", "Chỉ tài khoản đích thay đổi; ownership công ty không bị sai."),
        ("UIT-12", "[P1] Recruiter chưa kích hoạt hoặc token hết hạn.", "Tạo lại activation link; thử token cũ và token mới.", "Token cũ vô hiệu; token mới dùng một lần và được audit."),
    ]),
    ("7.3. Danh mục ngành nghề và kỹ năng", [
        ("UIT-13", "[P1] Có category/skill ACTIVE và ARCHIVED.", "Mở phân hệ; tìm kiếm; đổi bộ lọc trạng thái.", "Danh sách, usage count, key/slug, version và trạng thái chính xác."),
        ("UIT-14", "[P1] Tên/key category hợp lệ.", "Tạo category; đổi tên với expectedVersion; refresh.", "Tạo/cập nhật thành công; key chuẩn hóa; history/audit có actor."),
        ("UIT-15", "[P1] Category không được tin mở tham chiếu.", "Archive kèm lý do; thử dùng category đã archive cho tin mới.", "Archive thành công; tham chiếu inactive bị từ chối."),
        ("UIT-16", "[P0] Category/skill đang được tin mở sử dụng.", "Thử archive.", "Backend chặn và nêu ảnh hưởng; không làm hỏng tin hiện có."),
        ("UIT-17", "[P1] Category ARCHIVED.", "Reactivate kèm expectedVersion/lý do; thử tạo tin tham chiếu.", "Về ACTIVE và dùng lại được; version tăng."),
        ("UIT-18", "[P1] Dữ liệu skill hợp lệ và sai slug.", "Tạo, sửa, archive, reactivate skill; thử slug trùng/sai regex.", "Lifecycle giống category; validation/unique/conflict đúng; audit đủ."),
    ]),
    ("7.4. Báo cáo tuyển dụng", [
        ("UIT-19", "[P1] Dữ liệu nhiều khoa/ngành/company/status.", "Mở báo cáo; đối chiếu tổng, breakdown và danh sách.", "Tổng hợp khớp cùng tập lọc; pagination không lặp/mất dòng."),
        ("UIT-20", "[P1] Có dữ liệu trong/ngoài khoảng ngày.", "Kết hợp from/to, faculty, major, cohort, company, status, opportunityType và query.", "Bộ lọc giao nhau chính xác; reset trả danh sách ban đầu."),
        ("UIT-21", "[P0] Có giá trị bắt đầu bằng =,+,-,@ trong dữ liệu demo.", "Xuất CSV UTF-8 và mở bằng Excel/text editor.", "Không thực thi formula; cột/header/encoding/row count đúng; tên tệp hợp lệ."),
        ("UIT-22", "[P1] Cùng bộ lọc CSV.", "Xuất XLSX; mở workbook; đối chiếu ngẫu nhiên 5 dòng.", "Workbook mở không repair; header/date/number đúng; dữ liệu khớp CSV."),
        ("UIT-23", "[P0] Báo cáo >10.000 dòng hoặc request export sai role.", "Thử export quá giới hạn và bằng Student/Company.", "Giới hạn được áp dụng; sai role 403; mỗi export UIT được audit kèm filter/row count."),
    ]),
    ("7.5. Duyệt tin tuyển dụng", [
        ("UIT-24", "[P1] Có tin PENDING_UIT_REVIEW.", "Mở hàng đợi; chọn tin; đối chiếu company, deadline, category, mô tả, yêu cầu.", "Chi tiết đúng version; chỉ tin chờ duyệt xuất hiện."),
        ("UIT-25", "[P0] Tin đủ điều kiện.", "Chọn Phê duyệt & công khai; xác nhận.", "Tin thành RECRUITING; hiển thị cho Student; Company nhận thông báo; history/audit đủ."),
        ("UIT-26", "[P0] Tin thiếu nội dung.", "Chọn Yêu cầu chỉnh sửa; nhập lý do >=5 ký tự.", "Tin thành REVISION_REQUIRED; Company thấy đúng note và có thể sửa/gửi lại."),
        ("UIT-27", "[P0] Tin vi phạm quy định.", "Chọn Từ chối; nhập lý do; xác nhận.", "Tin thành REJECTED terminal; không công khai; Company nhận reason."),
        ("UIT-28", "[P0] Tin đã được xử lý hoặc hai admin cùng thao tác.", "Retry cùng command; gửi decision khác với version cũ.", "Cùng command không tạo history trùng; thao tác muộn bị từ chối."),
    ]),
    ("7.6. Xác minh tài liệu và duyệt hồ sơ", [
        ("UIT-29", "[P1] Có tài liệu PENDING/VERIFIED/REJECTED.", "Mở danh sách; lọc status; tìm MSSV/tên/tệp; chọn bản ghi.", "Danh sách và số lượng đúng; chi tiết có student, type, size, uploadedAt."),
        ("UIT-30", "[P0] Tài liệu PENDING.", "Mở PDF bằng CTA; chọn Xác minh hợp lệ; thêm note nội bộ.", "URL ký ngắn hạn; status VERIFIED; reviewer/time/note/audit đủ; không đổi lại."),
        ("UIT-31", "[P0] Tài liệu PENDING không hợp lệ.", "Chọn Từ chối; nhập lý do; xác nhận.", "Status REJECTED; Student nhận lý do; decision bất biến."),
        ("UIT-32", "[P1] Có đơn UIT_REVIEWING với snapshot.", "Mở hàng đợi; xem hồ sơ, eligibility, note và từng PDF snapshot.", "Dữ liệu khớp thời điểm nộp; không bị thay theo CV hiện tại."),
        ("UIT-33", "[P0] Đơn cần bổ sung.", "Chọn Yêu cầu bổ sung; chọn reasonCode, deadline và note.", "Đơn thành NEEDS_SUPPLEMENT; Student nhận task/deadline; history/audit/notification đủ."),
        ("UIT-34", "[P0] Đơn không đủ điều kiện.", "Chọn Từ chối; nhập note; xác nhận.", "Đơn thành UIT_REJECTED terminal; Student nhận thông báo."),
        ("UIT-35", "[P0] Đơn đủ điều kiện.", "Chọn Duyệt & chuyển doanh nghiệp; xác nhận modal.", "Đơn thành FORWARDED_TO_COMPANY; chỉ company sở hữu nhìn thấy; hai bên nhận notification."),
        ("UIT-36", "[P0] Đơn đã rời UIT_REVIEWING.", "Gửi supplement/reject/forward lần nữa.", "Transition bị từ chối; không tạo thêm history/notification."),
    ]),
    ("7.7. Xác nhận và theo dõi kỳ thực tập", [
        ("UIT-37", "[P1] Đơn ACCEPTED_PENDING_UIT_CONFIRMATION.", "Mở Theo dõi kết quả; chọn hồ sơ; mở PDF offer và đối chiếu startDate.", "Queue/detail/PDF đúng; URL ký; CTA confirm chỉ hiện khi đủ điều kiện."),
        ("UIT-38", "[P0] Sinh viên có một đơn accept và các đơn active khác.", "Chọn Xác nhận nơi thực tập; nhập ngày bắt đầu; xác nhận.", "Đơn chính HIRED; tạo placement; đơn khác WITHDRAWN/ACCEPTED_OTHER_JOB và lịch liên quan CANCELLED trong cùng transaction."),
        ("UIT-39", "[P1] Có placement HIRED/STARTED/COMPLETED.", "Lọc danh sách; mở chi tiết; xem version, effective date, actor history.", "Lifecycle tách khỏi application HIRED; timeline và availableActions chính xác."),
        ("UIT-40", "[P0] Placement HIRED; ngày bắt đầu hợp lệ.", "Chọn Ghi nhận bắt đầu; xác nhận.", "Placement thành STARTED; application vẫn HIRED; history/audit/notifications đủ."),
        ("UIT-41", "[P0] Placement STARTED; ngày hoàn thành hợp lệ.", "Chọn Ghi nhận hoàn thành; xác nhận.", "Placement thành COMPLETED; mở quyền gửi phiếu cho Student/Company."),
        ("UIT-42", "[P0] Ngày tương lai, complete trước start, expectedVersion cũ, retry cùng command.", "Gửi từng request sai và retry.", "Validation/state/conflict đúng; retry idempotent; không có history trùng."),
        ("UIT-43", "[P0] Placement COMPLETED có 0/1/2 phiếu.", "Mở chi tiết bằng UIT; so sánh khi Company/Student gửi.", "UIT xem được cả hai phiếu; không sửa; actor/time/ratings/text chính xác."),
    ]),
    ("7.8. Thông báo UIT", [
        ("UIT-44", "[P1] Có thông báo tin/hồ sơ/placement.", "Mở inbox; đánh dấu đã đọc/tất cả; kiểm tra unread count.", "Chỉ cập nhật thông báo của admin; count nhất quán."),
        ("UIT-45", "[P1] Notification có resourceType/resourceId.", "Bấm mở từ inbox.", "Đi đúng queue/detail và chọn đúng bản ghi; resource đã terminal hiển thị an toàn."),
    ]),
]


COMPANY_MODULES = [
    ("8.1. Tổng quan và hồ sơ doanh nghiệp", [
        ("COM-01", "[P1] Company có tin, ứng viên, lịch và hire.", "Mở dashboard; đối chiếu 4 metric, ứng viên cần xử lý và hiệu quả từng tin.", "Số liệu chỉ thuộc companyId đăng nhập và khớp API."),
        ("COM-02", "[P2] Company mới không có dữ liệu hoặc API lỗi.", "Mở dashboard; mô phỏng empty/loading/error; retry.", "Empty copy/CTA hợp lý; error không vỡ layout; retry phục hồi."),
        ("COM-03", "[P1] Company ACTIVE có profile.", "Mở Hồ sơ doanh nghiệp; đối chiếu tên, mã, lĩnh vực, quy mô, website, mô tả.", "Dữ liệu đúng company của recruiter; không hiển thị trường nội bộ nhạy cảm."),
        ("COM-04", "[P1] Dữ liệu profile hợp lệ và sai định dạng.", "Sửa profile; lưu; refresh; thử URL/chuỗi quá dài.", "Hợp lệ được lưu, version/audit tăng; validation chặn dữ liệu sai."),
        ("COM-05", "[P1] Profile vừa cập nhật.", "Đăng nhập Student; mở danh bạ đối tác và company detail.", "Thông tin công khai mới hiển thị; recruiter email/nội bộ không lộ."),
        ("COM-06", "[P0] Recruiter Company A biết ID Company B.", "Gọi profile/job/candidate/document/placement thuộc Company B.", "Trả 404 theo resource; dashboard/report không rò dữ liệu Company B."),
    ]),
    ("8.2. Tin tuyển dụng", [
        ("COM-07", "[P1] Company ACTIVE.", "Tạo tin với title, type, work mode, location, positions, deadline, description, requirements, benefits; Lưu nháp.", "Tin DRAFT xuất hiện đúng company; version=giá trị khởi tạo; audit có actor."),
        ("COM-08", "[P1] Tin DRAFT.", "Mở chỉnh sửa; thay nhiều trường; lưu; refresh.", "Nội dung/version đúng; danh sách/filter cập nhật."),
        ("COM-09", "[P0] Form thiếu trường bắt buộc, positions ngoài 1-1000, deadline không hợp lệ.", "Thử Lưu nháp và Lưu & gửi duyệt.", "CTA/validation chặn; backend không lưu dữ liệu không hợp lệ."),
        ("COM-10", "[P0] Tin DRAFT hợp lệ.", "Chọn Lưu & gửi UIT duyệt; double-click/retry.", "Tin PENDING_UIT_REVIEW; UIT nhận notification; cùng command không tạo history trùng."),
        ("COM-11", "[P0] Tin REVISION_REQUIRED có latestReview.note.", "Mở tin; đọc feedback; sửa; gửi lại.", "Tin trở về PENDING_UIT_REVIEW; feedback cũ vẫn trong history; UIT thấy version mới."),
        ("COM-12", "[P1] Có DRAFT/PENDING/REVISION/RECRUITING/REJECTED.", "Đổi bộ lọc trạng thái; tìm tin; mở tin editable.", "Danh sách đúng; chỉ DRAFT/REVISION_REQUIRED có CTA chỉnh sửa."),
        ("COM-13", "[P0] Tin đã PENDING/RECRUITING/REJECTED.", "Thử PATCH/submit trực tiếp bằng version cũ hoặc sai state.", "Backend từ chối; tin/history/notification không bị thay đổi ngoài ý muốn."),
    ]),
    ("8.3. Ứng viên, tài liệu và sàng lọc", [
        ("COM-14", "[P0] UIT đã chuyển hồ sơ cho Company A; có hồ sơ chưa chuyển.", "Mở board Ứng viên; lọc/tìm; so sánh hai hồ sơ.", "Chỉ hồ sơ FORWARDED_TO_COMPANY trở đi và thuộc Company A xuất hiện."),
        ("COM-15", "[P1] Đơn FORWARDED_TO_COMPANY.", "Chọn Bắt đầu xem; refresh board.", "Đơn thành COMPANY_REVIEWING; history/audit/notification theo contract; CTA thay đổi."),
        ("COM-16", "[P0] Đơn đã được UIT chuyển.", "Mở chi tiết; đối chiếu student, job, snapshot; mở từng PDF.", "Metadata đúng snapshot; URL ký ngắn hạn; không trả storage key; lần mở được audit."),
        ("COM-17", "[P0] Company A biết application/document ID của Company B.", "Gọi detail/download/start-review/reject.", "Trả 404; không lộ student, fileName, status hoặc company sở hữu."),
        ("COM-18", "[P0] Đơn FORWARDED_TO_COMPANY/COMPANY_REVIEWING.", "Chọn Không phù hợp; nhập note; xác nhận.", "Đơn NOT_SUITABLE terminal; Student nhận reason; không có CTA transition."),
        ("COM-19", "[P1] Board có nhiều job/status.", "Tìm tên/MSSV/vị trí; lọc job; quan sát 4 cột pipeline.", "Mỗi đơn xuất hiện một lần ở cột đúng; count khớp danh sách."),
    ]),
    ("8.4. Phỏng vấn, kết quả và offer", [
        ("COM-20", "[P0] Đơn COMPANY_REVIEWING.", "Chọn Mời PV; nhập scheduledAt, mode, location/meeting URL, interviewer; Gửi lời mời.", "Tạo interview; đơn INTERVIEW_INVITED; Student nhận thông báo/lịch."),
        ("COM-21", "[P1] Dữ liệu lịch thiếu/sai: quá khứ, online không URL, chuỗi quá dài.", "Gửi form từng trường hợp.", "Validation rõ ràng; không tạo interview hoặc transition dở dang."),
        ("COM-22", "[P1] Có lịch upcoming/history.", "Mở Lịch phỏng vấn; lọc scope; mở resource liên quan.", "Chỉ lịch của company; student confirmation/cancel/result được phản ánh."),
        ("COM-23", "[P0] Đơn INTERVIEW_INVITED.", "Chọn kết quả FAIL; nhập note/reasonCode.", "Đơn INTERVIEW_FAILED terminal; Student nhận kết quả; không có offer."),
        ("COM-24", "[P0] Đơn INTERVIEW_INVITED.", "Chọn Đạt & offer; nhập startDate/note; gửi không kèm PDF.", "Tạo RecruitmentResult PASS; đơn OFFER_PENDING_STUDENT; Student nhận notification."),
        ("COM-25", "[P0] PDF offer hợp lệ <10 MB.", "Tạo upload, PUT R2, complete; gửi PASS; mở lại PDF.", "Offer document gắn đúng application; ba role được phép mở bằng URL ký; storage key không lộ."),
        ("COM-26", "[P0] Offer PNG/EXE/>10 MB/MIME-signature sai hoặc application không thuộc company.", "Thử upload/complete/download.", "File sai bị từ chối; cross-company trả 404; không tạo recruitment result sai."),
        ("COM-27", "[P0] Đơn đã terminal hoặc offer đã gửi.", "Retry cùng command; gửi decision khác/key khác.", "Cùng command idempotent; transition sai bị chặn; không có offer/history trùng."),
    ]),
    ("8.5. Thông báo và đánh giá thực tập", [
        ("COM-28", "[P1] Có notification từ UIT/Student/system.", "Mở inbox; mark one/read-all; bấm deep link.", "Unread count đúng; chỉ thông báo company; deep link mở đúng candidate/job."),
        ("COM-29", "[P0] Company A biết notification ID của Company B.", "Gọi read notification của Company B.", "Trả 404; unread của cả hai không bị sai."),
        ("COM-30", "[P1] Placement COMPLETED thuộc company, chưa có phiếu Company.", "Mở candidate detail; chọn 4 ratings; nhập strengths/improvements; gửi.", "Tạo một phiếu bất biến; Student và UIT xem được phần Company."),
        ("COM-31", "[P0] Phiếu Company đã gửi hoặc placement chưa COMPLETED.", "Gửi lần hai hoặc gửi sớm; dùng Company B.", "Không tạo phiếu trùng/sớm/cross-company; Company không nhìn thấy nội dung phiếu riêng tư của Student."),
    ]),
]


CROSS_ROLE_CASES = [
    ("E2E-01", "[P0] Happy flow đầy đủ.", "Company tạo/gửi tin → UIT duyệt → Student nộp → UIT chuyển → Company PV/PASS/offer → Student nhận → UIT confirm.", "Trạng thái đúng từ DRAFT đến HIRED; history/audit/notification đủ; không bỏ qua role."),
    ("E2E-02", "[P0] Nhánh sửa tin.", "Company gửi tin → UIT yêu cầu sửa → Company xem note, chỉnh sửa, gửi lại → UIT duyệt.", "REVISION_REQUIRED được xử lý đúng; feedback cũ không mất; tin cuối RECRUITING."),
    ("E2E-03", "[P0] Nhánh bổ sung hồ sơ.", "Student nộp → UIT yêu cầu bổ sung → Student upload/resubmit → UIT chuyển.", "NEEDS_SUPPLEMENT và bản snapshot mới chính xác; Company chỉ nhận bản UIT đã chuyển."),
    ("E2E-04", "[P0] UIT từ chối hồ sơ.", "Student nộp → UIT reject → Student mở notification/đơn.", "UIT_REJECTED terminal; reason hiển thị đúng; Company không bao giờ nhìn thấy."),
    ("E2E-05", "[P0] Company chọn không phù hợp hoặc FAIL.", "UIT chuyển → Company reject; lặp lại checkpoint cho nhánh interview FAIL.", "NOT_SUITABLE/INTERVIEW_FAILED đúng reason; Student nhận thông báo; không có offer."),
    ("E2E-06", "[P0] Student rút đơn.", "Rút khi UIT_REVIEWING và khi COMPANY_REVIEWING.", "WITHDRAWN; actor Student; UIT/Company liên quan nhận thông báo; CTA phía Company bị khóa."),
    ("E2E-07", "[P0] Student hủy phỏng vấn.", "Company mời → Student xác nhận → Student chọn không tham gia.", "Interview CANCELLED và application WITHDRAWN trong cùng luồng; Company thấy kết quả."),
    ("E2E-08", "[P0] Student từ chối offer.", "Company PASS/offer → Student mở PDF → decline kèm reason.", "OFFER_DECLINED terminal; không tạo placement; Company nhận reason."),
    ("E2E-09", "[P0] Sinh viên có >=3 đơn active.", "Nhận một offer → UIT confirm placement.", "Đơn chính HIRED; đơn khác WITHDRAWN/ACCEPTED_OTHER_JOB; interview liên quan hủy; tất cả cùng transaction."),
    ("E2E-10", "[P0] Placement HIRED.", "UIT start → UIT complete → Company gửi phiếu → Student xem/gửi phiếu → UIT đối chiếu.", "Placement COMPLETED; mỗi role một phiếu; Company không thấy text Student; UIT xem cả hai."),
    ("E2E-11", "[P0] Hai company và hai student.", "Dùng ID chéo để mở đơn/tệp/lịch/offer/phiếu; thử endpoint sai role.", "Ownership trả 404; sai role 403; không có metadata nhạy cảm trong body/log."),
    ("E2E-12", "[P0] Hai tab cùng resource và network retry.", "Gửi cùng command, khác command, expectedVersion cũ; quan sát history/notification.", "Idempotency không nhân đôi side effect; optimistic lock chặn ghi đè; notification dedupe hoạt động."),
]


NONFUNCTIONAL_CASES = [
    ("NFT-01", "[P0] Không có token hoặc token sai issuer/audience/chữ ký/hết hạn.", "Gọi đại diện mọi nhóm endpoint bảo vệ.", "401 AUTH_ACCESS_TOKEN_MISSING/INVALID; không parse payload nghiệp vụ trước auth."),
    ("NFT-02", "[P0] Token hợp lệ của từng role.", "Chạy ma trận Student/UIT/Company trên các endpoint riêng.", "Role đúng pass; role sai 403; endpoint dùng chung hoạt động theo contract."),
    ("NFT-03", "[P0] Resource chéo student/company.", "Chạy negative ownership cho job, application, document, interview, offer, evaluation, notification.", "Tất cả trả 404; không lộ resource existence."),
    ("NFT-04", "[P0] Presigned URL hợp lệ.", "Kiểm tra response API, URL, expiry; thử mở sau hết hạn và tìm storage_key trong UI/log.", "URL sống ngắn; hết hạn bị từ chối; storage key/secret không bị lộ."),
    ("NFT-05", "[P0] Tập tệp biên PDF.", "Kiểm tra MIME, extension, size, magic bytes, zero-byte, corrupted PDF, tên tệp Unicode/rất dài.", "Chỉ PDF hợp lệ <=10 MB được complete; tên tệp không gây path traversal/XSS."),
    ("NFT-06", "[P0] Chuỗi bắt đầu =,+,-,@ trong trường báo cáo.", "Xuất CSV/XLSX; mở trong Excel.", "CSV formula injection được neutralize; XLSX không tự chạy công thức ngoài ý muốn."),
    ("NFT-07", "[P0] Hai request cập nhật cùng version.", "Gửi song song trên company/taxonomy/job/placement.", "Chỉ một request thành công; request còn lại nhận conflict và không ghi đè."),
    ("NFT-08", "[P0] Cùng Idempotency-Key và command khác nhau.", "Retry submit/decision/transition khi mạng ngắt sau response.", "Cùng key trả kết quả nhất quán; không nhân đôi history/audit/notification/file."),
    ("NFT-09", "[P1] Gọi API bảo vệ và trang frontend.", "Kiểm tra Cache-Control, content type, CORS, security headers và không cache response nhạy cảm.", "API /api/v1 trả private, no-store; origin cho phép đúng; cookie HttpOnly/Secure/SameSite theo môi trường."),
    ("NFT-10", "[P1] Desktop/tablet/mobile.", "Chạy luồng chính ở 1440x900, 768x1024 và 390x844; zoom 200%.", "Không che CTA/modal/table; navigation dùng được; text không bị cắt."),
    ("NFT-11", "[P1] Chỉ dùng bàn phím và screen reader smoke.", "Tab qua login, nav, filter, table row, dialog; Esc/Đóng; kiểm tra label/role/alert/focus.", "Thứ tự focus hợp lý; focus không thoát modal; input/icon/alert có tên truy cập."),
    ("NFT-12", "[P1] API chậm, 4xx/5xx và danh sách rỗng.", "Mô phỏng trên mỗi portal ít nhất một danh sách và một form.", "Loading/empty/error/retry/success rõ; double submit bị khóa; form không mất dữ liệu vô lý."),
    ("NFT-13", "[P1] Checkpoint có dữ liệu đủ lớn.", "Đo p95 login/dashboard/list/filter/transition và export trong môi trường test ổn định.", "Ghi baseline; không timeout; query/pagination không tăng tuyến tính bất thường. Ngưỡng cuối do nhóm chốt trước release."),
    ("NFT-14", "[P0] Backend và database đang chạy/ngắt.", "Gọi /health và /health/database; tắt DB tạm thời rồi khôi phục.", "Health phản ánh đúng; API trả lỗi có trace ID; không trả stack trace/secret."),
    ("NFT-15", "[P0] Có backup test và môi trường restore cô lập.", "Restore; chạy db check và happy flow smoke trên bản phục hồi.", "Schema/data/reference integrity đúng; ba role đăng nhập và đọc dữ liệu được."),
    ("NFT-16", "[P1] Có pending events và scheduler secret.", "Chạy daily pending đúng/sai secret; mô phỏng email provider lỗi và retry.", "Sai secret bị từ chối; outbox không rollback nghiệp vụ; retry/dedupe không gửi trùng."),
    ("NFT-17", "[P0] Chạy các transition quan trọng.", "Kiểm tra structured log, trace ID, audit và state history; tìm token/password/storage key trong log.", "Actor/action/target/time/reason/version đủ; log không chứa secret, token, presigned URL đầy đủ hoặc file nhạy cảm."),
]


FULL_FLOW_STEPS = [
    ("1", "Chuẩn bị", "Reset checkpoint demo/E2E; mở 3 profile trình duyệt riêng.", "Ba phiên tách cookie; dữ liệu đúng checkpoint."),
    ("2", "Company", "Đăng nhập VNG; mở Tin tuyển dụng; tạo tin 'Kiểm thử 3 vai trò - <timestamp>'.", "Tin DRAFT hoặc được lưu và gửi thẳng PENDING_UIT_REVIEW."),
    ("3", "UIT", "Mở Duyệt tin; chọn tin; phê duyệt công khai.", "Tin RECRUITING; Company có notification; tin xuất hiện cho Student."),
    ("4", "Student", "Mở Việc làm; tìm tin; xem chi tiết; bắt đầu ứng tuyển.", "Chi tiết khớp tin đã duyệt; flow chia sẻ tài liệu hiển thị."),
    ("5", "Student", "Chọn CV/tài liệu; tick consent; xác nhận và gửi.", "Tạo một đơn UIT_REVIEWING và snapshot; UIT nhận task."),
    ("6", "UIT", "Mở Duyệt hồ sơ; xem snapshot/PDF; chọn Duyệt & chuyển doanh nghiệp.", "Đơn FORWARDED_TO_COMPANY; Company thấy đúng hồ sơ."),
    ("7", "Company", "Mở Ứng viên; tìm theo tin; chọn Bắt đầu xem.", "Đơn COMPANY_REVIEWING; Student thấy cập nhật."),
    ("8", "Company", "Chọn Mời PV; nhập lịch trong tương lai, mode và link/venue.", "Tạo lịch; đơn INTERVIEW_INVITED; Student nhận notification."),
    ("9", "Student", "Mở Lịch phỏng vấn; xác nhận tham gia; mở link nếu online.", "Lịch CONFIRMED; Company thấy trạng thái mới."),
    ("10", "Company", "Chọn Đạt & offer; nhập startDate; upload PDF offer hợp lệ; gửi.", "PASS + OFFER_PENDING_STUDENT; PDF private mở được bằng URL ký."),
    ("11", "Student", "Mở Đơn ứng tuyển; mở PDF; chọn Nhận offer.", "ACCEPTED_PENDING_UIT_CONFIRMATION; các đơn khác chưa bị đóng."),
    ("12", "UIT", "Mở Theo dõi kết quả; đối chiếu offer; xác nhận nơi thực tập.", "Đơn chính HIRED; tạo placement HIRED; đơn active khác WITHDRAWN."),
    ("13", "Student + Company", "Làm mới đơn, candidate board và notification ở các bên liên quan.", "Một đơn HIRED; các đơn khác có reason ACCEPTED_OTHER_JOB; không có CTA sai."),
    ("14", "UIT", "Trong Vòng đời thực tập, ghi nhận bắt đầu với ngày hợp lệ.", "Placement STARTED; application vẫn HIRED."),
    ("15", "UIT", "Ghi nhận hoàn thành sau ngày bắt đầu.", "Placement COMPLETED; mở form phiếu cho Company/Student."),
    ("16", "Company", "Mở candidate detail; gửi 4 ratings, strengths và improvements.", "Tạo phiếu COMPANY một lần; Student và UIT xem được."),
    ("17", "Student", "Mở đơn HIRED; xem phiếu Company; gửi phiếu Student.", "Tạo phiếu STUDENT một lần; text riêng tư không hiển thị cho Company."),
    ("18", "UIT", "Mở placement COMPLETED; đối chiếu hai phiếu và history.", "UIT xem đủ hai phiếu, actor/time/ratings; toàn bộ luồng có audit/notification."),
]


BRANCH_MATRIX = [
    ("Tin cần sửa", "PENDING_UIT_REVIEW", "UIT yêu cầu chỉnh sửa", "REVISION_REQUIRED", "Company thấy feedback; chỉ sửa/gửi lại được."),
    ("Tin bị từ chối", "PENDING_UIT_REVIEW", "UIT từ chối", "REJECTED", "Terminal; không công khai."),
    ("Hồ sơ cần bổ sung", "UIT_REVIEWING", "UIT yêu cầu bổ sung", "NEEDS_SUPPLEMENT", "Student resubmit về UIT_REVIEWING."),
    ("UIT từ chối hồ sơ", "UIT_REVIEWING", "UIT reject", "UIT_REJECTED", "Terminal; Company không thấy."),
    ("Company loại sớm", "FORWARDED/REVIEWING", "Không phù hợp", "NOT_SUITABLE", "Student nhận reason."),
    ("Phỏng vấn không đạt", "INTERVIEW_INVITED", "Company ghi FAIL", "INTERVIEW_FAILED", "Không tạo offer."),
    ("Student rút", "Trạng thái cho phép", "Rút đơn", "WITHDRAWN", "Khóa CTA; thông báo bên liên quan."),
    ("Student hủy PV", "INTERVIEW_INVITED", "Không thể tham gia", "WITHDRAWN", "Interview CANCELLED cùng luồng."),
    ("Student từ chối offer", "OFFER_PENDING_STUDENT", "Từ chối", "OFFER_DECLINED", "Không có placement."),
    ("UIT chốt nơi thực tập", "ACCEPTED_PENDING_UIT_CONFIRMATION", "Confirm placement", "HIRED", "Các đơn active khác WITHDRAWN/ACCEPTED_OTHER_JOB."),
]


def total_module_cases(modules):
    return sum(len(cases) for _, cases in modules)


def add_role_section(doc, heading, intro, modules):
    doc.add_page_break()
    doc.add_heading(heading, level=1)
    doc.add_paragraph(intro)
    for index, (title, cases) in enumerate(modules):
        add_test_table(doc, title, cases, page_before=index > 0)


def build_document():
    doc = Document()
    configure_styles(doc)
    set_document_furniture(doc)
    bullet_num = create_numbering(doc, fmt="bullet", marker="•", bullet_font=FONT)
    decimal_num = create_numbering(doc, fmt="decimal", marker="%1.")

    # Cover: editorial_cover pattern, reduced to a monochrome academic treatment.
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(110)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("TRƯỜNG ĐẠI HỌC CÔNG NGHỆ THÔNG TIN")
    set_run_font(r, 11, bold=True)
    p = doc.add_paragraph("QUY TRÌNH KIỂM THỬ 3 VAI TRÒ", style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph("UIT CAREER HUB", style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        set_run_font(run, 28, bold=True)
    p = doc.add_paragraph(
        "Sổ tay kiểm thử thủ công, nghiệp vụ liên vai trò và kiểm tra an toàn",
        style="Subtitle",
    )
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(40)
    add_simple_table(
        doc,
        ["Phạm vi", "Căn cứ", "Cập nhật"],
        [(
            "Sinh viên - UIT Admin - Doanh nghiệp",
            "Giao diện, Java API, state machine, RBAC và Playwright hiện có",
            "21/08/2026",
        )],
        [2700, 4500, 2160],
        center_columns=(0, 2),
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(48)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Nhóm sinh viên: Phạm Thị Kiều Diễm - 23520286  |  Nguyễn Trương Ngọc Hân - 23520434")
    set_run_font(r, 10.5, bold=True)

    doc.add_page_break()
    doc.add_heading("1. Mục đích và phạm vi", level=1)
    doc.add_paragraph(
        "Tài liệu này là quy trình kiểm thử dùng trước khi demo, bàn giao hoặc bảo vệ đồ án. "
        "Mục tiêu là xác nhận toàn bộ chức năng hiện có của ba portal cùng các bất biến quan trọng ở backend."
    )
    for item in (
        "Kiểm thử theo vai trò: Sinh viên, UIT Admin và Doanh nghiệp.",
        "Kiểm thử luồng liên vai trò từ tạo tin đến hoàn thành thực tập và gửi đánh giá.",
        "Kiểm thử nhánh ngoại lệ, RBAC, ownership, idempotency, optimistic locking, tệp private và báo cáo.",
        "Ghi bằng chứng đủ để truy vết defect và tái hiện lần chạy.",
    ):
        add_bullet(doc, item, bullet_num)
    add_note(
        doc,
        "Giới hạn: ",
        "Tài liệu bám theo chức năng đang có trong frontend và Java backend. Các hướng phát triển như UIT SSO, FCM, saved jobs, chat realtime, AI và mobile app không được coi là test bắt buộc của phiên bản hiện tại.",
    )

    doc.add_heading("2. Cấu trúc tài liệu", level=1)
    for item in (
        "Mục 3-5: chuẩn bị môi trường, dữ liệu và cách ghi kết quả.",
        "Mục 6: test portal Sinh viên.",
        "Mục 7: test portal UIT Admin.",
        "Mục 8: test portal Doanh nghiệp.",
        "Mục 9: happy flow và các nhánh E2E liên vai trò.",
        "Mục 10-12: phi chức năng, lệnh test tự động và biên bản lần chạy.",
    ):
        add_bullet(doc, item, bullet_num)

    doc.add_heading("3. Chuẩn bị môi trường", level=1)
    doc.add_heading("3.1. Nguyên tắc an toàn", level=2)
    for item in (
        "Chỉ reset database demo/E2E cô lập; không trỏ DATABASE_URL vào production.",
        "Không trình chiếu .env, access token, refresh cookie, R2 secret hoặc presigned URL đầy đủ.",
        "Dùng ba Chrome profile/incognito context riêng để cookie của các vai trò không ghi đè nhau.",
        "Mỗi lần chạy ghi commit, URL, database/checkpoint, trình duyệt và thời gian.",
    ):
        add_bullet(doc, item, bullet_num)

    doc.add_heading("3.2. Tài khoản demo", level=2)
    add_simple_table(
        doc,
        ["Vai trò", "Email", "Mật khẩu", "Mục đích"],
        [
            ("UIT Admin", "admin.career@uit.edu.vn", "Admin@12345", "Duyệt tin/hồ sơ, xác nhận placement."),
            ("Sinh viên chính", "20521067@student.uit.edu.vn", "Student@12345", "Luồng nhiều đơn và nhận offer."),
            ("Sinh viên phụ", "21520881@student.uit.edu.vn", "Student@12345", "Upload/PDF offer và ownership chéo."),
            ("VNG recruiter", "recruiter@vng.example", "Company@12345", "Luồng Company chính."),
            ("FPT recruiter", "recruiter@fpt.example", "Company@12345", "Kiểm tra cross-company và đơn tự đóng."),
        ],
        [1500, 3060, 1860, 2940],
    )
    add_note(doc, "Cảnh báo: ", "Các mật khẩu trên chỉ dùng cho demo/test. Phải tắt hoặc xoay trước khi dùng môi trường thật.")

    doc.add_heading("3.3. Trình tự khởi động", level=2)
    steps = (
        "Tại thư mục gốc, bật cờ reset an toàn: $env:ALLOW_DEMO_RESET = \"true\".",
        "Chạy pnpm db:demo:reset; sau khi xong đặt ALLOW_DEMO_RESET về false.",
        "Terminal 1 chạy pnpm dev:backend; chờ health và database sẵn sàng.",
        "Terminal 2 chạy pnpm dev:frontend; mở URL Vite được in ra.",
        "Mở ba profile trình duyệt, đăng nhập ba role và chạy AUTH-01..03 trước.",
    )
    for step in steps:
        add_numbered(doc, step, decimal_num)

    doc.add_heading("4. Chiến lược thực thi và mức bao phủ", level=1)
    counts = [
        ("Dùng chung / xác thực", len(AUTH_CASES), "Login, refresh, logout, activation, revoked token, RBAC."),
        ("Sinh viên", total_module_cases(STUDENT_MODULES), "7 màn hình và lifecycle đơn/placement."),
        ("UIT Admin", total_module_cases(UIT_MODULES), "9 phân hệ quản trị, duyệt và báo cáo."),
        ("Doanh nghiệp", total_module_cases(COMPANY_MODULES), "6 màn hình, tin, candidate pipeline, offer, phiếu."),
        ("Liên vai trò", len(CROSS_ROLE_CASES), "Happy flow, nhánh lỗi và transaction nhiều đơn."),
        ("Phi chức năng", len(NONFUNCTIONAL_CASES), "Bảo mật, file, concurrency, a11y, responsive, restore."),
    ]
    add_simple_table(doc, ["Nhóm", "Số test", "Phạm vi"], counts, [2100, 1080, 6180], center_columns=(1,))
    total = sum(row[1] for row in counts)
    add_labelled(doc, "Tổng số kịch bản: ", f"{total}. Cần pass toàn bộ P0 và không còn defect P1 cản trở buổi demo.")
    add_status_legend(doc)

    doc.add_page_break()
    doc.add_heading("5. Test dùng chung và xác thực", level=1)
    add_test_table(doc, "5.1. Đăng nhập, phiên và phân quyền", AUTH_CASES)

    add_role_section(
        doc,
        "6. Quy trình test vai trò Sinh viên",
        "Thực thi trên tài khoản sinh viên chính, sau đó dùng sinh viên phụ cho negative ownership và dữ liệu offer độc lập.",
        STUDENT_MODULES,
    )
    add_role_section(
        doc,
        "7. Quy trình test vai trò UIT Admin",
        "UIT là vai trò trung gian kiểm soát chất lượng. Khi test transition cần kiểm tra đồng thời UI, response API, state history, audit và notification.",
        UIT_MODULES,
    )
    add_role_section(
        doc,
        "8. Quy trình test vai trò Doanh nghiệp",
        "Dùng VNG cho luồng chính và FPT cho ownership/cross-company. Mọi candidate, tệp và dashboard phải bị giới hạn bởi companyId trong token.",
        COMPANY_MODULES,
    )

    doc.add_page_break()
    doc.add_heading("9. Luồng kiểm thử liên vai trò", level=1)
    doc.add_heading("9.1. Happy flow đầy đủ", level=2)
    doc.add_paragraph(
        "Luồng này là gate quan trọng nhất trước demo. Không sửa trực tiếp status bằng SQL; mọi transition phải đi qua UI/API thật."
    )
    add_chunked_matrix(
        doc,
        ["Bước", "Vai trò", "Thao tác", "Checkpoint mong đợi"],
        FULL_FLOW_STEPS,
        [720, 1320, 3720, 3600],
        chunk_size=6,
        font_size=8.9,
        center_columns=(0, 1),
        continuation_title="9.1. Happy flow đầy đủ",
    )
    doc.add_page_break()
    doc.add_heading("9.2. Ma trận nhánh ngoại lệ", level=2)
    add_chunked_matrix(
        doc,
        ["Nhánh", "Từ trạng thái", "Hành động", "Đến trạng thái", "Bất biến cần kiểm tra"],
        BRANCH_MATRIX,
        [1500, 1740, 1740, 1740, 2640],
        chunk_size=5,
        font_size=8.7,
        center_columns=(1, 3),
        continuation_title="9.2. Ma trận nhánh ngoại lệ",
    )
    add_test_table(doc, "9.3. Test case E2E", CROSS_ROLE_CASES, page_before=True)

    doc.add_page_break()
    doc.add_heading("10. Kiểm thử phi chức năng và an toàn", level=1)
    doc.add_paragraph(
        "Các kịch bản này được thực thi sau khi happy flow ổn định. Negative test nên chạy bằng API client tự động để quan sát chính xác HTTP status và error code."
    )
    add_test_table(doc, "10.1. Security, reliability, usability và recovery", NONFUNCTIONAL_CASES)

    doc.add_page_break()
    doc.add_heading("11. Bộ test tự động và cổng chất lượng", level=1)
    add_simple_table(
        doc,
        ["Lệnh", "Phạm vi", "Kết quả bắt buộc"],
        [
            ("pnpm typecheck", "Java compile/type safety và contract biên dịch.", "Exit code 0."),
            ("pnpm test", "Frontend test + backend JUnit/Spring Test.", "Tất cả test pass; không test bị skip ngoài danh sách cho phép."),
            ("pnpm build", "Vite production build + Java backend build.", "Build thành công; không lỗi bundling/compile."),
            ("pnpm e2e:smoke", "Đăng nhập ba role và tải dashboard API.", "Ba portal pass trên database E2E cô lập."),
            ("pnpm e2e:full", "Happy flow, revision, supplement, reject, withdraw, cancel interview, decline, ownership, taxonomy, report.", "Toàn bộ @full pass; không flake khi chạy lại."),
            ("pnpm smoke:offer-flow", "HTTP + Neon + R2 cho PDF offer; chỉ khi đã bật cờ demo mutation.", "Upload/download/offer flow pass; reset checkpoint sau khi chạy."),
        ],
        [2160, 4260, 2940],
    )
    add_note(
        doc,
        "Thiết lập E2E: ",
        "Cần NODE_ENV=test, DATABASE_URL_E2E trỏ database có tên chứa 'e2e' và ALLOW_E2E_RESET=true. Không dùng database runtime/production. Khi test fail, lưu playwright-report, screenshot, video và trace.zip.",
    )

    doc.add_heading("12. Ghi nhận kết quả và tiêu chí kết thúc", level=1)
    for item in (
        "Tất cả test P0 PASS; không có lỗ hổng lộ dữ liệu, sai role, sai ownership hoặc transition không hợp lệ.",
        "Happy flow E2E pass liên tiếp ít nhất hai lần trên checkpoint sạch; không có history/notification/file trùng.",
        "Không có defect P1 chặn demo; defect P2 còn lại có owner, mức ảnh hưởng và kế hoạch xử lý.",
        "Build, unit/integration test, smoke và full E2E pass trên commit dùng để bàn giao.",
        "Evidence đủ để giảng viên/người review tái hiện: input, actor, timestamp, response/status, screenshot và defect ID nếu có.",
    ):
        add_bullet(doc, item, bullet_num)
    add_test_run_record(doc)

    doc.core_properties.title = "Quy trình kiểm thử 3 vai trò - UIT Career Hub"
    doc.core_properties.subject = "Manual test guide, E2E workflow and security coverage"
    doc.core_properties.author = "Nhóm sinh viên 23520286 - 23520434"
    doc.core_properties.keywords = "UIT Career Hub, test plan, Student, UIT Admin, Company, E2E, RBAC"
    doc.save(OUTPUT)
    print(f"Created: {OUTPUT}")
    print(f"Test cases: {total}")
    print(f"Size: {OUTPUT.stat().st_size} bytes")


if __name__ == "__main__":
    build_document()
